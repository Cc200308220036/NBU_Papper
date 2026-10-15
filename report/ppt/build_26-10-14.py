"""Build the editable 2026-10-14 NBU_Papper research readout.

Source deck: Simple Light Mode reference.pptx (masters/theme retained).
Run: /home/cyw/anaconda3/envs/hilserl_sim/bin/python report/ppt/build_26-10-14.py
"""
from pathlib import Path
import copy

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.util import Inches, Pt
from pptx.oxml.xmlchemy import OxmlElement

ROOT = Path(__file__).resolve().parents[2]
REF = Path('/home/cyw/.codex/plugins/cache/openai-curated-remote/openai-templates/0.1.1/skills/artifact-template-simple-light-mode/assets/reference.pptx')
OUT = Path(__file__).with_name('26-10-14.pptx')
IMG = ROOT / 'semantic_pallet_planner/outputs/EXP_1C_DEEPSEEK_FLASH_FINAL_V1/selector_inputs/vlm_visual/DC_00354/workspace.png'

prs = Presentation(REF)
# Keep the reference masters/layouts/theme; replace its lorem-ipsum sample slides.
for sid in list(prs.slides._sldIdLst):
    prs.part.drop_rel(sid.rId)
    prs.slides._sldIdLst.remove(sid)
blank = prs.slide_layouts[7]
W, H = 13.333, 7.5
C = {
    'ink': RGBColor(25, 31, 37),
    'muted': RGBColor(91, 100, 109),
    'line': RGBColor(217, 222, 226),
    'panel': RGBColor(247, 248, 249),
    'soft': RGBColor(236, 241, 242),
    'accent': RGBColor(25, 105, 113),
    'accent_light': RGBColor(222, 239, 239),
    'red': RGBColor(166, 67, 59),
    'red_light': RGBColor(250, 235, 232),
    'green': RGBColor(34, 119, 83),
    'green_light': RGBColor(228, 244, 234),
    'white': RGBColor(255, 255, 255),
}
FONT = 'Noto Sans CJK SC'


def inch(v): return Inches(v)


def font_ea(run):
    run.font.name = FONT
    rpr = run._r.get_or_add_rPr()
    ea = rpr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ea')
    if ea is None:
        ea = OxmlElement('a:ea')
        rpr.append(ea)
    ea.set('typeface', FONT)


def textbox(slide, x, y, w, h, text, size=18, bold=False, color='ink',
            align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0, linespace=1.13):
    shp = slide.shapes.add_textbox(inch(x), inch(y), inch(w), inch(h))
    tf = shp.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = inch(margin)
    tf.margin_top = tf.margin_bottom = inch(margin)
    tf.vertical_anchor = valign
    for idx, line in enumerate(str(text).split('\n')):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(0)
        p.line_spacing = linespace
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = C[color]
        font_ea(run)
    return shp


def rect(slide, x, y, w, h, fill='panel', line=None, radius=False):
    typ = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(typ, inch(x), inch(y), inch(w), inch(h))
    s.fill.solid(); s.fill.fore_color.rgb = C[fill]
    if line:
        s.line.color.rgb = C[line]
        s.line.width = Pt(0.7)
    else:
        s.line.fill.background()
    if radius:
        try: s.adjustments[0] = 0.06
        except Exception: pass
    return s


def line(slide, x1, y1, x2, y2, color='line', width=1):
    s = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, inch(x1), inch(y1), inch(x2), inch(y2))
    s.line.color.rgb = C[color]; s.line.width = Pt(width)
    return s


def footer(slide, page, source):
    line(slide, .48, 6.86, 12.85, 6.86, width=.65)
    textbox(slide, .5, 6.93, 7.3, .20, 'NBU_Papper  ·  开发验证数据，非正式新 test', 9.2, color='muted')
    textbox(slide, 8.0, 6.93, 4.4, .20, source, 8.6, color='muted', align=PP_ALIGN.RIGHT)
    textbox(slide, 12.5, 6.92, .35, .22, f'{page:02d}', 9.2, color='muted', align=PP_ALIGN.RIGHT)


def new_slide(page, title, subtitle='', source=''):
    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, W, H, fill='white')
    textbox(s, .5, .36, 12.2, .62, title, 28, True)
    if subtitle:
        textbox(s, .52, 1.05, 12.1, .49, subtitle, 14.2, color='muted')
    footer(s, page, source)
    return s


def card(slide, x, y, w, h, label, big, detail='', tone='accent'):
    rect(slide, x, y, w, h, fill='panel', line='line', radius=True)
    textbox(slide, x+.22, y+.18, w-.44, .35, label, 13.2, color='muted')
    textbox(slide, x+.22, y+.62, w-.44, .72, big, 30, True, tone)
    if detail:
        textbox(slide, x+.22, y+1.43, w-.44, h-1.57, detail, 12.4, color='muted')


def note(slide, text):
    ns = slide.notes_slide
    tf = ns.notes_text_frame
    if tf is not None:
        tf.text = text


def pill(slide, x, y, w, text, fill='soft', color='accent'):
    rect(slide, x, y, w, .33, fill=fill, radius=True)
    textbox(slide, x+.07, y+.03, w-.14, .24, text, 10.8, True, color)


# 1. Cover
s = prs.slides.add_slide(blank)
rect(s, 0, 0, W, H, fill='white')
textbox(s, .51, .43, 6.8, .4, 'NBU_PAPPER  /  阶段研究汇报', 13.4, True, color='muted')
line(s, .51, 1.39, 2.65, 1.39, 'accent', 4)
textbox(s, .5, 1.76, 11.85, 1.93, '属性感知码垛\n近期进展与核心实验', 40, True, linespace=1.16)
textbox(s, .53, 4.12, 10.9, 1.0, '从固定候选选择走向“自然语言政策解析 + 确定性执行 + 在线安全闭环”', 20, color='muted')
textbox(s, .53, 6.57, 6, .35, '2026-10-14   |   V2 · O2a · O1b · O2b · O3 计划', 13, color='muted')
note(s, '开场说明：本次主要讲最近几个阶段的原始实验成果，区分工程检查通过与研究假设得到支持。当前 O2b 的真实调用矩阵完成，O3 在线阶段尚未实现。')

# 2. Progress map
s = new_slide(2, '进展总览：O2b 工程矩阵已跑通，O3 是下一步',
              '近期交付沿着“指标 → 协议 → 数据 → 决策 → 在线布局”逐层推进。', '来源：各阶段 acceptance_report.json')
stages = [
    ('1C / V1','固定候选平台','360 题 · 2160 行','PASS'),
    ('V2','属性指标','60 题 · Direct 真调用','开发完成'),
    ('O2a','协议修复','新 Prompt 20/20 首答合法','人工签认待办'),
    ('O1b','语言/政策数据','30 来源订单 · 180 题','自动 PASS'),
    ('O2b','AST 混合决策','1080/1080 回放','工程 PASS'),
    ('O3','在线属性布局','实际轨迹分叉与完成率','待开发'),
]
for i,(name,topic,value,status) in enumerate(stages):
    x=.5+i*2.12
    rect(s,x,2.05,1.96,3.4,fill='white' if i==5 else 'panel',line='line',radius=True)
    rect(s,x,2.05,1.96,.07,fill='accent' if i<5 else 'line')
    textbox(s,x+.18,2.35,1.6,.37,name,17,True,color='accent' if i<5 else 'muted')
    textbox(s,x+.18,2.96,1.6,.52,topic,16,True)
    textbox(s,x+.18,3.68,1.6,.75,value,12.7,color='muted')
    textbox(s,x+.18,4.75,1.6,.39,status,11,True,color='accent' if i<5 else 'muted')
textbox(s,.53,5.81,12.2,.52,'判断边界：O2b 工程 PASS ≠ Flash 相对强规则有优势；O3 只有方案，尚无在线实验结果。',16,True)
note(s,'V1 360 题六方法 2160 行；V2 60 题、30 个状态但仅 10 独立来源订单。O2a 第二批 20 道开发题新 Prompt 首答合法 20/20；O1b 30 独立 train 来源订单形成 180 题。O2b 1080 行回放通过，但完整统计设计待完成。')

# 3. Data chain and historical image
s = new_slide(3, '框架数据链路：模型解释业务，程序负责关系和安全',
              '所有非特权方法共享公开输入与同一物理候选；隐藏真值只用于评测。', '来源：attribute_o2b.py / O1b public_requests')
flow = [
    ('公开状态与业务指令','当前箱、已有箱、属性；buffer_size=1'),
    ('物理候选生成','冻结生成器；Top-K=16；只给 Cxx 编号'),
    ('决策接口','Direct 选编号  /  文本→受限 AST→执行器'),
    ('提交屏障与复核','保留原始提议；拦截业务违规；独立物理检查'),
]
for i,(a,b) in enumerate(flow):
    y=1.65+i*1.12
    rect(s,.52,y,6.65,.85,fill='panel',line='line',radius=True)
    textbox(s,.74,y+.12,2.35,.27,a,14.8,True)
    textbox(s,3.1,y+.12,3.77,.43,b,12.2,color='muted')
    if i<3: textbox(s,3.57,y+.86,.5,.24,'↓',14,True,color='accent',align=PP_ALIGN.CENTER)
s.shapes.add_picture(str(IMG),inch(7.51),inch(1.77),width=inch(5.25))
textbox(s,7.56,4.88,5.0,.5,'阶段 1C 历史状态渲染示意（DC_00354）',11,color='muted')
rect(s,7.51,5.45,5.25,.91,fill='accent_light',radius=True)
textbox(s,7.72,5.61,4.84,.58,'state_hash 绑定物理状态\npolicy_digest 区分同状态不同政策',13.2,True,color='accent')
note(s,'输入链路来自 O1b public_request 和 O2b runner。图像是阶段 1C 的历史状态渲染，仅用于说明码垛场景；不是 O2b 的图像输入，也不能当成视觉增量证据。模型/解析器只读公开当前状态、属性、指令和候选关系；隐藏 AST、Oracle 分数、未来箱体、真实候选 ID 映射留在评测侧。')

# 4. O1b
s = new_slide(4, 'O1b：题目数量之外，必须看独立来源与政策机会',
              '原 train 来源派生开发集；同物理状态的六题只改变公开政策。', '来源：O1b acceptance_report / independent_verification')
card(s,.52,1.74,2.9,2.21,'独立来源订单','30','与原 validation、V2/O2a、test 来源零重合')
card(s,3.66,1.74,2.9,2.21,'固定候选题','180','A/B/C 三轨 × 双政策；共享状态与候选')
card(s,6.80,1.74,2.9,2.21,'Top-K 业务可行','168/180','另 12 题全枚举也无业务满足解')
card(s,9.94,1.74,2.9,2.21,'真正切换机会','10/30','仅 10 个来源订单出现最优动作切换')
rect(s,.52,4.39,12.32,1.69,fill='panel',line='line',radius=True)
for i,(h,t) in enumerate([('A 规范表达','已支持语法回归'),('B 改写表达','相同政策不同表述'),('C 条件分支','按当前箱属性选择政策')]):
    x=.8+i*4.05
    textbox(s,x,4.67,3.7,.38,h,17,True)
    textbox(s,x,5.15,3.7,.42,t,13,color='muted')
    if i<2: line(s,x+3.78,4.62,x+3.78,5.71)
textbox(s,.54,6.27,12.2,.35,'注意：180 题 ≠ 180 个独立订单；无解题必须单列 abstain，不能计作业务成功。',14,True,color='red')
note(s,'O1b 对 90 个候选状态进行冻结条件筛查，保留 30 个状态；每个来源订单一个状态，A/B/C 各两条政策，总计 180 题和 90 个配对。Top-K 有解 168 题，全枚举无业务可行解 12 题。真正因政策变化产生最优集合差异的 30 对来自 10 个独立来源订单。语义属性为 synthetic。')

# 5. O2a
s = new_slide(5, 'O2a：协议合法率改善，但不能替代业务判断',
              '新 Prompt 修复输出结构；候选内无解时仍需确定性业务屏障。', '来源：O2a _02 acceptance_report / metrics')
rect(s,.53,1.62,6.0,3.69,fill='panel',line='line',radius=True)
textbox(s,.81,1.88,5.42,.42,'首答响应合法率（20 道同题开发题）',17,True)
for i,(name,n,clr) in enumerate([('旧 Prompt',18,'muted'),('新 Prompt',20,'accent')]):
    y=2.67+i*1.06
    textbox(s,.83,y,1.55,.39,name,15)
    rect(s,2.53,y+.04,3.19,.27,fill='line')
    rect(s,2.53,y+.04,3.19*n/20,.27,fill=clr)
    textbox(s,5.63,y-.02,.70,.44,f'{n}/20',13,True,align=PP_ALIGN.RIGHT)
rect(s,6.79,1.62,6.04,3.69,fill='panel',line='line',radius=True)
textbox(s,7.05,1.88,5.48,.42,'适用硬约束满足（8 道题）',17,True)
for i,(name,n,clr) in enumerate([('旧 Prompt',5,'muted'),('新 Prompt',6,'accent'),('Public Parser',6,'accent')]):
    y=2.55+i*.72
    textbox(s,7.05,y,2.04,.32,name,13.7)
    rect(s,9.2,y+.05,2.76,.21,fill='line')
    rect(s,9.2,y+.05,2.76*n/8,.21,fill=clr)
    textbox(s,12.01,y-.03,.54,.34,f'{n}/8',13.5,True,align=PP_ALIGN.RIGHT)
rect(s,.53,5.62,12.3,.78,fill='red_light',radius=True)
textbox(s,.76,5.83,11.85,.38,'新 Prompt 在 2 道业务无解题仍误提交；20/20 是开发集观察，不保证未来请求永远 ≥98%。',14.6,True,color='red')
note(s,'V2 的 8/60 个首答无效包括 6 个多余 type 字段和 2 个 JSON 语法错误。O2a 严格限制七字段、request_id/state_hash/policy_digest 三项绑定和至多一次格式修复。_02 使用另一批 10 来源订单的 20 道双政策题，旧 Prompt 首答 18/20、新 Prompt 20/20；适用硬约束满足 5/8 对 6/8，公开 Parser 也是 6/8。两道无解题误提交说明协议修复不是语义可靠性的全部。')

# 6. O2b method comparison
s = new_slide(6, 'O2b 六方法同题对照：混合架构可靠，但未超强规则',
              '原始提议统计；业务屏障后的提交结果另计。', '来源：O2b metrics/summary.json')
headers=['方法','可避免硬违规','最优等价候选','输入权限 / 解释']
rows=[
 ('Geometry-Greedy','27/111  ·  24.3%','81/168  ·  48.2%','公开几何'),
 ('固定属性启发式','3/111  ·  2.7%','129/168  ·  76.8%','公开属性'),
 ('Flash Direct-Text','21/111  ·  18.9%','93/167  ·  55.7%','公开事实，直接选编号'),
 ('Public Parser＋Executor','0/111','168/168  ·  100%','公开自然语言＋确定性执行'),
 ('Flash AST＋Executor','0/111','168/168  ·  100%','公开自然语言→AST＋同一执行器'),
 ('真值 AST 参考','0/111','168/168  ·  100%','隐藏真值；特权参考'),
]
x0=.53; widths=[3.25,2.8,2.95,3.25]; xs=[x0]
for w in widths: xs.append(xs[-1]+w)
rect(s,.53,1.69,12.26,.58,fill='soft')
for i,h in enumerate(headers): textbox(s,xs[i]+.1,1.82,widths[i]-.18,.29,h,12.8,True)
for ri,row in enumerate(rows):
    y=2.28+ri*.63
    rect(s,.53,y,12.26,.62,fill='accent_light' if ri in (3,4) else ('panel' if ri%2==0 else 'white'))
    for ci,val in enumerate(row):
        textbox(s,xs[ci]+.1,y+.13,widths[ci]-.16,.4,val,13.2,ri in (3,4),color='accent' if ri in (3,4) else 'ink')
textbox(s,.55,6.27,12.2,.35,'分母：111 = 硬规则适用且 Top-K 有解；168 = 候选内有解；Direct 的有效原始提议分母为 167。',12.7,color='muted')
note(s,'六方法同题、同物理候选。原始可避免业务硬违规 Geometry 27/111，固定属性 3/111，Direct 21/111，Public Parser、Flash AST、真值 AST 均 0/111。最优等价候选率分别 81/168、129/168、93/167、168/168、168/168、168/168。真值 AST 是特权上界。Flash AST 超过 Direct 但与同权限 Public Parser 完全持平，所以不支持 VLM 相对强规则的增量优势。')

# 7. Concrete case
s = new_slide(7, '逐题案例：C08 新增禁止接触，C10 满足业务要求',
              'O1B_GROUP_004_A0：易碎承载≤1 kg；电子品不得与重箱直接面接触；再比较重箱低位/中心。', '来源：O2b public_request / per_decision.csv')
rect(s,.54,1.86,5.9,3.91,fill='red_light',line='line',radius=True)
pill(s,.82,2.12,2.2,'Flash Direct  →  C08',fill='white',color='red')
textbox(s,.83,2.75,5.3,.6,'新增禁止面接触',21,True,color='red')
textbox(s,.83,3.43,5.2,.73,'item_04（heavy）↔ item_05（electronics）',16,True)
textbox(s,.83,4.43,5.2,.58,'原始提议：新增业务违规 1',15.2)
textbox(s,.83,5.02,5.2,.48,'提交结果：屏障拦截 → abstain',15.2,True,color='red')
rect(s,6.79,1.86,5.99,3.91,fill='green_light',line='line',radius=True)
pill(s,7.07,2.12,3.35,'Flash AST / Public Parser  →  C10',fill='white',color='green')
textbox(s,7.08,2.75,5.3,.6,'没有这对新增禁止接触',20.2,True,color='green')
textbox(s,7.08,3.43,5.25,.88,'相同公开候选事实，由执行器按硬约束与优先级排序。',15.4)
textbox(s,7.08,4.43,5.25,.58,'原始提议：新增业务违规 0',15.2)
textbox(s,7.08,5.02,5.25,.48,'提交结果：C10；命中最优等价集合',15.2,True,color='green')
textbox(s,.55,6.11,12.2,.41,'读图关键：屏障阻止了 C08 的错误提交，但不能将其算作 Direct 自身“选择正确”。',14,True)
note(s,'这个例子来自 O2b 的公开请求和 per_decision.csv。C08 的 public_relations.delta.new_contact_pairs 中含 item_04 与 item_05 的接触；item_04 为 heavy，item_05 为 electronics。C10 无该新增接触。Direct 的 proposed_label=C08，screen_intervened=True，committed_label 为空；Flash AST 和 Public Parser 均提出并提交 C10。案例用于解释机制，不可凭单例作总体结论。')

# 8. Evidence and boundaries
s = new_slide(8, '证据边界：工程检查通过，论文优势仍待新任务验证',
              '技术复查、统计设计、在线验证和视觉贡献是不同层次。', '来源：O2b replay/acceptance；Agent review')
card(s,.54,1.71,2.91,2.12,'独立回放','1080/1080','O2b 六方法计划记录完整')
card(s,3.69,1.71,2.91,2.12,'已提交物理违规','0','不代表原始业务提议零违规',tone='green')
card(s,6.84,1.71,2.91,2.12,'Flash AST 严格一致','80/180','结构一致率；选址结果不能代替解析准确率',tone='red')
card(s,9.99,1.71,2.84,2.12,'政策切换来源','10/30','有效机会来自 10 个独立订单',tone='red')
rect(s,.54,4.31,12.29,1.75,fill='panel',line='line',radius=True)
textbox(s,.84,4.56,11.7,.37,'当前不能声称的结果',17.3,True)
textbox(s,.84,5.04,11.62,.76,'① Flash 优于强 Public Parser   ② 完整订单布局改善   ③ 图像带来增量   ④ 新 test 上的统计结论',14.7,color='muted')
textbox(s,.55,6.27,12.15,.32,'O2a/O1b 有 Agent 独立技术复查；两名真人的人工签认仍为 PENDING。',12.7,color='muted')
note(s,'O2b replay_report PASS，1080/1080，失败列表为空，已提交物理违规 0；acceptance_report 中 offline/remote engineering PASS，但 o2b_full_acceptance=PENDING_STATISTICAL_DESIGN，research_status=DEVELOPMENT_EVIDENCE_ONLY。Flash AST 严格 AST 结构精确率 80/180；不应由 100% 候选最优率反推 AST 每字段都对。Agent 复核 12/12 O2a 错例与 180/180 O1b AST，但并非两位真人独立签认。')

# 9. Research value
s = new_slide(9, '论文研究方向：找到“模型必要”的信息和任务边界',
              '创新要由同权限强对照和完整布局结果支撑，而不是先设定模型胜出。', '来源：O3 路线 / O2b 强基线结果')
ideas=[
 ('01','语言政策解释','新来源、未见表述/组合/例外；对比充分开发的 Public Parser。','证明“语言适应”是否真有增量'),
 ('02','在线属性布局','每步按本方法真实状态再生成候选；比较属性任务成功与完成率。','证明单步收益是否能累积'),
 ('03','视觉属性缺口','仅在图像确有独有属性时测试；对比 OCR 和真值 Text 上界。','分清感知贡献与规划贡献'),
]
for i,(num,head,body,bottom) in enumerate(ideas):
    x=.55+i*4.24
    rect(s,x,1.8,3.97,4.43,fill='panel',line='line',radius=True)
    textbox(s,x+.27,2.13,3.45,.44,num,15,True,color='accent')
    textbox(s,x+.27,2.78,3.45,.57,head,21,True)
    textbox(s,x+.27,3.67,3.42,1.37,body,15.1,color='muted')
    line(s,x+.27,5.26,x+3.67,5.26)
    textbox(s,x+.27,5.44,3.43,.58,bottom,12.7,True,color='accent')
note(s,'当前实验只有文本和真值 JSON，因此不应称视觉 VLM 贡献。强 Public Parser 已在本批受限语法上达到参考上界，下一轮必须在新来源、新语言表达和政策变化上与其同权限比较。O3 在线任务用完整 episode 指标检验长期布局；O4 才可能研究图像独有属性。若规则始终相当，应明确 VLM 不必要的子域与混合架构适用边界。')

# 10. O3 plan
s = new_slide(10, '下一阶段：先完成 O3 在线闭环，再冻结新 test',
              '先把框架接好并验证安全性；真实模型调用和研究结论随后分阶段推进。', '来源：O3 在线码垛框架与论文创新验证路线')
steps=[
 ('O3a · 接线','在线状态、policy_digest、AST、关系特征、业务屏障与独立回放。'),
 ('O3a · 离线验收','几何/固定属性/Public Parser/混合接口/真值在线参考；终止记录完整，物理违规 0。'),
 ('O3b · Validation','真实 Flash 在线调用；报告 Episode Attribute Success、完成率、属性原始量与屏障干预。'),
 ('新 test · 冻结后执行','先用独立开发来源确定组数、重复次数与阈值；再做正式配对矩阵。'),
]
for i,(h,b) in enumerate(steps):
    y=1.76+i*1.08
    pill(s,.54,y+.12,1.12,f'{i+1:02d}',fill='accent_light')
    textbox(s,1.87,y+.03,3.21,.39,h,16.5,True)
    textbox(s,5.00,y+.03,7.6,.71,b,14.2,color='muted')
    if i<3: line(s,.96,y+.45,.96,y+1.1,'accent',1.15)
rect(s,.54,6.23,12.28,.46,fill='soft',radius=True)
textbox(s,.76,6.31,11.82,.30,'近期实验不记录、不估算、不考核费用；保留请求、错误、版本和必要时延。O3 目前尚未实现。',12.4,True,color='accent')
note(s,'O3a 的工程验收不以 Flash 胜出为门槛。既有 PalletEnvironment 和 EpisodeRunner 使用早期协议，需新增版本化在线 runner/adapter，将 O2b 的公开政策解析、属性评价和业务屏障接到实际逐步提交状态。各方法从相同订单开始，但后续根据自己的状态重生成候选。O3b 才在新独立 validation 来源上调用模型。正式新 test 的主要终点、样本量、重复次数和阈值在看 test 前冻结。用户目前不要求费用记录。')

prs.core_properties.title = '属性感知码垛：近期进展与核心实验'
prs.core_properties.subject = 'NBU_Papper V2/O2a/O1b/O2b/O3 research readout'
prs.core_properties.author = 'NBU_Papper project'
prs.core_properties.keywords = 'palletizing, AST, DeepSeek Flash, O2b, O3'
prs.save(OUT)
print(OUT)
print('slides', len(prs.slides), 'bytes', OUT.stat().st_size)
