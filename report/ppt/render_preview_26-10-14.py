"""Approximate raster preview for structural visual review of the editable PPTX.
Not a PowerPoint/LibreOffice render; PowerPoint remains the authoritative viewer.
"""
from pathlib import Path
from io import BytesIO
import subprocess
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN

HERE=Path(__file__).resolve().parent
PPT=HERE/'26-10-14.pptx'
OUT=HERE/'26-10-14_preview.png'
prs=Presentation(PPT)
FONT=subprocess.check_output(['fc-match','-f','%{file}','Noto Sans CJK SC'],text=True).strip()
W,H=1067,600
scale=W/prs.slide_width

def px(e): return int(e*scale)

def wrap(draw,text,font,width):
    chunks=[]
    for para in text.split('\n'):
        if not para: chunks.append('');continue
        line=''
        for c in para:
            if line and draw.textlength(line+c,font=font)>width:
                chunks.append(line);line=c
            else:line+=c
        if line:chunks.append(line)
    return chunks

def rgb(color, default=(24,30,36)):
    try:
        x=color.rgb
        return (x[0],x[1],x[2]) if x else default
    except Exception:return default

ims=[]
for slide in prs.slides:
    im=Image.new('RGB',(W,H),'white')
    draw=ImageDraw.Draw(im)
    for sh in slide.shapes:
        x,y,w,h=map(px,(sh.left,sh.top,sh.width,sh.height))
        if sh.shape_type==MSO_SHAPE_TYPE.PICTURE:
            pic=Image.open(BytesIO(sh.image.blob)).convert('RGB')
            pic.thumbnail((w,h),Image.Resampling.LANCZOS)
            im.paste(pic,(x,y))
        elif sh.shape_type==MSO_SHAPE_TYPE.LINE:
            draw.line((x,y,x+w,y+h),fill=rgb(sh.line.color,(215,220,225)),width=max(1,int((sh.line.width.pt if sh.line.width else 1)*W/960)))
        elif sh.shape_type==MSO_SHAPE_TYPE.AUTO_SHAPE:
            try:
                fill=rgb(sh.fill.fore_color,(255,255,255)) if sh.fill.type else None
            except Exception: fill=None
            if fill:
                draw.rounded_rectangle((x,y,x+w,y+h),radius=7,fill=fill)
        if sh.has_text_frame and sh.text.strip():
            tf=sh.text_frame
            cy=y+2
            for p in tf.paragraphs:
                if not p.runs: continue
                run=p.runs[0]
                size=run.font.size.pt if run.font.size else 16
                font=ImageFont.truetype(FONT,max(7,round(size*W/960)))
                ls=round(size*W/960*1.14)
                color=rgb(run.font.color)
                chunks=wrap(draw,p.text.replace('\v','\n'),font,max(5,w-4))
                for chunk in chunks:
                    if p.alignment==PP_ALIGN.RIGHT:
                        tx=x+w-2-draw.textlength(chunk,font=font)
                    elif p.alignment==PP_ALIGN.CENTER:
                        tx=x+(w-draw.textlength(chunk,font=font))/2
                    else:tx=x+2
                    draw.text((tx,cy),chunk,font=font,fill=color)
                    cy+=ls
    ims.append(im)
cols=2; rows=(len(ims)+1)//2
thumb_w=960;thumb_h=540;gap=20
sheet=Image.new('RGB',(cols*thumb_w+(cols+1)*gap,rows*thumb_h+(rows+1)*gap),(235,238,240))
for i,im in enumerate(ims):
    im=im.resize((thumb_w,thumb_h),Image.Resampling.LANCZOS)
    xx=gap+(i%cols)*(thumb_w+gap);yy=gap+(i//cols)*(thumb_h+gap)
    sheet.paste(im,(xx,yy))
    ImageDraw.Draw(sheet).rectangle((xx,yy,xx+thumb_w-1,yy+thumb_h-1),outline=(185,191,197),width=1)
sheet.save(OUT)
print(OUT, sheet.size)
