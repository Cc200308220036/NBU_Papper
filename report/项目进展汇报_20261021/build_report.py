"""Build the 2026-10-21 evidence-based progress report; never run an experiment/API.

Run with /home/cyw/anaconda3/envs/hilserl_sim/bin/python.
All writes are restricted to this report directory. Sources remain read-only.
"""
from pathlib import Path
from datetime import datetime, timezone, timedelta
import csv
import hashlib
import json
import math
import shutil
import textwrap
import zipfile
from io import BytesIO

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.patches import Rectangle
from mpl_toolkits.mplot3d.art3d import Poly3DCollection
import numpy as np
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
ASSETS = HERE/'assets'
ASSETS.mkdir(exist_ok=True)
EXP = ROOT/'semantic_pallet_planner/outputs/EXP_OPT_O3_LAYOUT_FLASH_V2_VALIDATION_20261020_01'
OLD = ROOT/'semantic_pallet_planner/outputs/EXP_OPT_O3B_BUSINESS_OBSERVE_20261016_REMOTE_01'
DIAG = ROOT/'report/analysis/O3_FAILURE_ATTRIBUTION_20261019_02'
FONT = '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'
BOLD = '/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc'
font_manager.fontManager.addfont(FONT)
plt.rcParams.update({'font.family':'Noto Sans CJK JP', 'axes.unicode_minus':False,
                     'font.size':12, 'axes.spines.top':False, 'axes.spines.right':False,
                     'figure.facecolor':'white','savefig.facecolor':'white'})
C = {'ink':'172D43','muted':'607487','navy':'102C42','teal':'007F82',
     'blue':'3673AC','amber':'D69531','red':'B6544F','white':'FFFFFF',
     'light':'F2F6F8','line':'D9E3E9','pale':'E5F2F0','warm':'FFF3DF'}
VARIANTS = ['original','relaxed_geometry','relaxed_wo_floor']
LABELS = ['原策略','预算＋几何','预算＋布局（去地面项）']
COLORS = ['#607487','#3673AC','#007F82']
SOURCES = {}

def read(path):
    path = Path(path)
    SOURCES[str(path.relative_to(ROOT))] = hashlib.sha256(path.read_bytes()).hexdigest()
    return json.loads(path.read_text(encoding='utf-8-sig'))

def rows(path):
    path = Path(path)
    SOURCES[str(path.relative_to(ROOT))] = hashlib.sha256(path.read_bytes()).hexdigest()
    with path.open(encoding='utf-8-sig',newline='') as f:
        return list(csv.DictReader(f))

acceptance = read(EXP/'acceptance_report.json')
replay = read(EXP/'replay_report.json')
usage = read(EXP/'api_usage.json')
summary = read(EXP/'metrics/summary.json')
episodes = rows(EXP/'metrics/per_episode.csv')
agree = rows(EXP/'metrics/parser_flash_agreement.csv')
execution = read(EXP/'audit/execution_checks.json')
paired = read(EXP/'metrics/paired_cluster_comparisons.json')
historical = read(OLD/'metrics/summary.json')
interventions = rows(DIAG/'single_action_interventions.csv')
divergences = read(DIAG/'divergence_details.json')
runtime = read(ROOT/'semantic_pallet_planner/outputs/EXP_OPT_O3_LAYOUT_STABLE_V2_PY311_20261020_01/cross_runtime_report.json')
latest_review = read(ROOT/'report/analysis/O3_LAYOUT_FLASH_V2_RESULTS_20261021_01/review_data.json')
assert acceptance['status'] == 'PASS' and acceptance['episodes'] == 240
assert replay['status'] == 'PASS' and replay['steps'] == 3042
assert usage['actual_attempts'] == 1502
assert len(agree) == 1502 and all(all(r[k]=='True' for k in ['ast_equal','same_state_action_equal','closed_loop_states_equal']) for r in agree)
aggregate_checks=0
for key,d in summary.items():
    mode,v,parser=key.split('/')
    rr=[r for r in episodes if r['mode']==mode and r['variant']==v and r['parser']==parser]
    assert len(rr)==d['episodes']
    assert sum(r['completed']=='True' for r in rr)==d['completed']
    assert sum(r['clean_completed']=='True' for r in rr)==d['clean_completed']
    assert sum(int(r['placed_items']) for r in rr)==d['placed_steps']
    assert sum(int(r['violating_steps']) for r in rr)==d['violating_steps']
    assert sum(int(r['final_business_violation_count']) for r in rr)==d['final_business_violation_instances']
    assert abs(sum(float(r['placed_item_ratio']) for r in rr)/len(rr)-d['mean_placed_item_ratio'])<1e-12
    aggregate_checks+=7
for mode in ['safe','observe_business']:
    for v in VARIANTS:
        assert summary[f'{mode}/{v}/flash'] == summary[f'{mode}/{v}/public_parser']

def savefig(fig,name):
    p = ASSETS/name
    fig.savefig(p,dpi=190,bbox_inches='tight')
    plt.close(fig)
    return p

def chart_historical():
    fig,ax = plt.subplots(figsize=(11.7,3.7))
    names=['geometry_greedy','fixed_attribute','public_parser','flash_policy_parser','truth_ast_executor']
    x=np.arange(5); w=.32
    ax.bar(x-w/2,[historical[n]['completed'] for n in names],w,color='#B8C8D6',label='完整放完')
    ax.bar(x+w/2,[historical[n]['attribute_success'] for n in names],w,color='#007F82',label='无违规完整放完')
    for i,n in enumerate(names):
        ax.text(i-w/2,historical[n]['completed']+.35,str(historical[n]['completed']),ha='center',fontsize=13)
        ax.text(i+w/2,historical[n]['attribute_success']+.35,str(historical[n]['attribute_success']),ha='center',fontsize=13)
    ax.set_xticks(x,['Geometry','固定属性','Public Parser','Flash AST','真值 AST'])
    ax.set_ylim(0,23);ax.set_ylabel('任务数 / 每方法20个');ax.legend(loc='upper right',frameon=False,ncol=2)
    ax.grid(axis='y',alpha=.18);ax.set_axisbelow(True)
    return savefig(fig,'01_历史五方法结果.png')

def chart_results(mode):
    ss=[summary[f'{mode}/{v}/public_parser'] for v in VARIANTS]
    fig,axes=plt.subplots(1,2,figsize=(12,3.8),gridspec_kw={'width_ratios':[1.2,1]})
    ax=axes[0];x=np.arange(3);w=.33
    ax.bar(x-w/2,[s['completed'] for s in ss],w,color='#B8C8D6',label='完整放完')
    ax.bar(x+w/2,[s['clean_completed'] for s in ss],w,color='#007F82',label='无违规完整放完')
    for i,s in enumerate(ss):
        ax.text(i-w/2,s['completed']+.35,str(s['completed']),ha='center',fontsize=13)
        ax.text(i+w/2,s['clean_completed']+.35,str(s['clean_completed']),ha='center',fontsize=13)
    ax.set_xticks(x,['原策略','预算＋几何','预算＋布局']);ax.set_ylim(0,23);ax.set_ylabel('任务数 / 每行20个')
    ax.legend(frameon=False,ncol=2,fontsize=10,loc='upper left');ax.grid(axis='y',alpha=.15);ax.set_axisbelow(True)
    ax=axes[1]
    if mode=='safe':
        vals=[100*s['mean_placed_item_ratio'] for s in ss]
        ax.barh(np.arange(3),vals,color=COLORS,height=.55)
        ax.set_yticks(np.arange(3),['原策略','预算＋几何','预算＋布局']);ax.invert_yaxis();ax.set_xlim(0,110)
        for i,v in enumerate(vals):ax.text(v+1,i,f'{v:.2f}%',va='center',fontsize=12)
        ax.set_xlabel('平均件数放置比例（%）')
    else:
        vals=[s['final_business_violation_instances'] for s in ss]
        ax.bar(x,vals,color=COLORS,width=.53)
        ax.set_xticks(x,['原策略','预算＋几何','预算＋布局']);ax.set_ylim(0,36);ax.set_ylabel('最终业务违规实例总数')
        for i,v in enumerate(vals):ax.text(i,v+.5,str(v),ha='center',fontsize=14)
    fig.tight_layout(w_pad=3)
    return savefig(fig,'02_安全模式.png' if mode=='safe' else '03_业务观察模式.png')

def chart_intervention():
    fig,ax=plt.subplots(figsize=(11.6,3.5))
    vals=[int(r['placed_delta']) for r in interventions]
    labels=[r['episode_id'].replace('O3_SCN_','').replace('_',' / ') for r in interventions]
    x=np.arange(len(vals));ax.bar(x,vals,color=['#007F82' if n>0 else '#B8C8D6' for n in vals],width=.6)
    for i,v in enumerate(vals):ax.text(i,v+.15,f'+{v}',ha='center',fontsize=12)
    ax.set_xticks(x,labels,fontsize=10);ax.set_ylabel('替换首次分叉动作后，新增放置箱数');ax.set_ylim(0,8.5)
    ax.grid(axis='y',alpha=.15);ax.set_axisbelow(True)
    return savefig(fig,'04_单动作干预.png')

def chart_heatmap():
    fig,ax=plt.subplots(figsize=(12,3.35))
    tasks=[f'O3_SCN_{i:04d}_P{p}' for i in range(51,61) for p in range(2)]
    mat=[]
    for v in VARIANTS:
        lookup={r['episode_id']:r for r in episodes if r['mode']=='observe_business' and r['variant']==v and r['parser']=='public_parser'}
        mat.append([1 if lookup[t]['clean_completed']=='True' else 0 for t in tasks])
    from matplotlib.colors import ListedColormap
    ax.imshow(mat,cmap=ListedColormap(['#E3E9EE','#007F82']),vmin=0,vmax=1,aspect='auto')
    ax.set_yticks(range(3),['原策略','预算＋几何','预算＋布局'])
    ax.set_xticks(range(20),[f'{i:04d}\nP{p}' for i in range(51,61) for p in range(2)],fontsize=9)
    for x in range(1,20,2):ax.axvline(x+.5,color='white',lw=3)
    for y in range(3):
        for x in range(20):ax.text(x,y,'✓' if mat[y][x] else '—',ha='center',va='center',color='white' if mat[y][x] else '#8998A5',fontsize=14)
    ax.set_title('绿色 = 无违规完整放完；灰色 = 该主指标未达成',fontsize=12,pad=15)
    return savefig(fig,'05_逐任务成功分布.png')

def chart_decomposition():
    vv=['relaxed_geometry','relaxed_wo_floor']
    by={v:{r['episode_id']:r for r in episodes if r['mode']=='observe_business' and r['variant']==v and r['parser']=='public_parser'} for v in vv}
    done={v:{t for t,r in by[v].items() if r['completed']=='True'} for v in vv}
    assert done[vv[0]]==done[vv[1]] and len(done[vv[0]])==18
    completed_viol=[];failed_viol=[]
    for v in vv:
        completed_viol.append(sum(int(r['final_business_violation_count']) for t,r in by[v].items() if t in done[v]))
        failed_viol.append(sum(int(r['final_business_violation_count']) for t,r in by[v].items() if t not in done[v]))
        assert sum(int(r['placed_items']) for t,r in by[v].items() if t in done[v])==250
    assert completed_viol==[25,17] and failed_viol==[5,0]
    fig,ax=plt.subplots(figsize=(10.8,3.4));x=np.arange(2)
    ax.barh(x,completed_viol,color='#007F82',height=.48,label='相同18个完成任务（各250箱）')
    ax.barh(x,failed_viol,left=completed_viol,color='#D69531',height=.48,label='0060双政策失败任务')
    for i,(a,b) in enumerate(zip(completed_viol,failed_viol)):
        ax.text(a/2,i,str(a),color='white',fontsize=18,ha='center',va='center')
        if b:ax.text(a+b/2,i,str(b),color='white',fontsize=18,ha='center',va='center')
        ax.text(a+b+.5,i,f'合计 {a+b}',fontsize=13,va='center')
    ax.set_yticks(x,['预算＋几何','预算＋布局']);ax.invert_yaxis();ax.set_xlim(0,36);ax.set_xlabel('最终业务违规实例数')
    ax.legend(frameon=False,ncol=2,loc='upper center',bbox_to_anchor=(.5,1.24),fontsize=11)
    ax.grid(axis='x',alpha=.15);ax.set_axisbelow(True);fig.tight_layout()
    return savefig(fig,'10_违规改善来源拆分.png')

def event_log(mode,variant,sid,policy='P1'):
    p=EXP/f'episodes/{mode}/{variant}/public_parser/O3_{sid}_{policy}/steps.jsonl'
    SOURCES[str(p.relative_to(ROOT))]=hashlib.sha256(p.read_bytes()).hexdigest()
    return [json.loads(s) for s in p.read_text().splitlines()]

def episode_row(mode,variant,sid,policy='P1'):
    return next(r for r in episodes if r['mode']==mode and r['variant']==variant and r['scenario_id']==sid and r['policy_id']==policy and r['parser']=='public_parser')

def draw_layout(ax,placed,title):
    for p in placed:
        l,w,h=p['oriented_size_mm'];pose=p['pose'];x=pose['x_mm']-l/2;y=pose['y_mm']-w/2;z=pose['z_base_mm']
        verts=np.array([[x,y,z],[x+l,y,z],[x+l,y+w,z],[x,y+w,z],[x,y,z+h],[x+l,y,z+h],[x+l,y+w,z+h],[x,y+w,z+h]])
        faces=[[verts[k] for k in inds] for inds in [[0,1,2,3],[4,5,6,7],[0,1,5,4],[1,2,6,5],[2,3,7,6],[3,0,4,7]]]
        color='#D69531' if p['fragile'] else ('#267B96' if p['weight_class']=='heavy' else '#9DB1C0')
        poly=Poly3DCollection(faces,facecolors=color,edgecolors='#354D5D',linewidths=.45,alpha=.92)
        ax.add_collection3d(poly)
    ax.plot([-600,600,600,-600,-600],[-500,-500,500,500,-500],[0]*5,color='#415B70',lw=1.2)
    ax.set(xlim=(-600,600),ylim=(-500,500),zlim=(0,850))
    ax.set_box_aspect((1200,1000,850));ax.view_init(24,-58)
    ax.set_xticks([-600,0,600]);ax.set_yticks([-500,0,500]);ax.set_zticks([0,400,850])
    ax.tick_params(labelsize=8,pad=0);ax.set_xlabel('X / mm',fontsize=9,labelpad=-1);ax.set_ylabel('Y / mm',fontsize=9,labelpad=-1)
    ax.set_title(title,fontsize=14,pad=4)
    for axis in [ax.xaxis,ax.yaxis,ax.zaxis]:axis.pane.fill=False

def chart_case(sid,policy='P1'):
    fig=plt.figure(figsize=(12,4.2))
    for i,v in enumerate(VARIANTS):
        es=event_log('observe_business',v,sid,policy);r=episode_row('observe_business',v,sid,policy)
        ax=fig.add_subplot(1,3,i+1,projection='3d')
        draw_layout(ax,es[-1]['after_placed'],f'{LABELS[i]}\n放置 {r["placed_items"]}/{r["planned_items"]}  |  最终违规 {r["final_business_violation_count"]}')
    fig.subplots_adjust(left=0,right=1,bottom=.04,top=.83,wspace=.01)
    return savefig(fig,f'06_{sid}_{policy}_真实布局.png')

def chart_divergence():
    # Plot actual saved post-action layouts from the diagnostic's source run.
    detail=next(d for d in divergences if d['episode_id']=='O3_SCN_0052_P0')
    fig,axes=plt.subplots(1,2,figsize=(10,4))
    for ax,method,title in zip(axes,['public_parser','geometry_greedy'],['Parser / Flash：偏向居中','Geometry：保留连续空间']):
        p=OLD/f'episodes/{method}/O3_SCN_0052_P0/steps.jsonl'
        SOURCES[str(p.relative_to(ROOT))]=hashlib.sha256(p.read_bytes()).hexdigest()
        e=json.loads(p.read_text().splitlines()[1])
        placed=e.get('after',{}).get('placed_items')
        if placed is None: placed=e.get('after_placed')
        if placed is None: raise ValueError(f'Unexpected historical layout fields: {list(e)}')
        for j,b in enumerate(placed):
            l,w,h=b['oriented_size_mm'];x=b['pose']['x_mm']-l/2;y=b['pose']['y_mm']-w/2
            ax.add_patch(Rectangle((x,y),l,w,facecolor=['#B5C8D7','#78AFB1'][j],edgecolor='#39546A',lw=1.5))
            ax.text(x+l/2,y+w/2,b['item_id'].replace('item_','箱 '),ha='center',va='center',fontsize=13)
        surf=detail[method+'_after_surface'];x0,x1,y0,y1=surf['largest_floor_clear_rectangle']
        ax.add_patch(Rectangle((x0,y0),x1-x0,y1-y0,fill=False,edgecolor='#B6544F',lw=2,ls='--'))
        ax.set(xlim=(-600,600),ylim=(-500,500),aspect='equal');ax.set_title(title,fontsize=15)
        ax.set_xlabel(f'最大连续矩形 {surf["largest_floor_clear_rectangle_mm2"]/1e6:.4f} m²',fontsize=13)
        ax.set_xticks([-600,0,600]);ax.set_yticks([-500,0,500]);ax.grid(alpha=.1)
    fig.tight_layout(w_pad=4)
    return savefig(fig,'07_0052首次分叉.png')

historical_chart=chart_historical()
safe_chart=chart_results('safe')
observe_chart=chart_results('observe_business')
intervention_chart=chart_intervention()
heatmap=chart_heatmap()
decomposition=chart_decomposition()
case54=chart_case('SCN_0054')
case60=chart_case('SCN_0060')
divergence_chart=chart_divergence()
for source,name in [(DIAG/'figures/O3_SCN_0054_P1.png','08_历史0054失败归因.png'),
 (ROOT/'report/visualization/O3_完整订单_SCN0054_P0_26-10-17/final_layouts.png','09_历史五方法完整订单.png')]:
    SOURCES[str(source.relative_to(ROOT))]=hashlib.sha256(source.read_bytes()).hexdigest()
    shutil.copyfile(source,ASSETS/name)

prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
SLIDES=[]

def rgb(c):return RGBColor.from_string(C.get(c,c))
def rect(slide,x,y,w,h,fill='light',line=None):
    s=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    s.fill.solid();s.fill.fore_color.rgb=rgb(fill)
    if line:s.line.color.rgb=rgb(line)
    else:s.line.fill.background()
    return s

def tx(slide,x,y,w,h,value,size=19,color='ink',bold=False,align=PP_ALIGN.LEFT):
    s=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=s.text_frame
    tf.clear();tf.word_wrap=True;tf.auto_size=MSO_AUTO_SIZE.NONE
    tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(str(value).split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.alignment=align;p.line_spacing=1.13;p.space_after=Pt(5)
        r=p.add_run();r.text=line;r.font.name='Noto Sans CJK SC';r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=rgb(color)
        ea=OxmlElement('a:ea');ea.set('typeface','Noto Sans CJK SC');r._r.get_or_add_rPr().append(ea)
    return s

def picture(slide,path,x,y,w,h):
    iw,ih=Image.open(path).size;scale=min(w/iw,h/ih);ww,hh=iw*scale,ih*scale
    return slide.shapes.add_picture(str(path),Inches(x+(w-ww)/2),Inches(y+(h-hh)/2),width=Inches(ww),height=Inches(hh))

def base(title,subtitle,source,notes):
    s=prs.slides.add_slide(prs.slide_layouts[6]);i=len(prs.slides)
    rect(s,0,0,13.334,.12,'teal');tx(s,.58,.38,12.15,.6,title,28,bold=True)
    tx(s,.6,1.04,12.1,.55,subtitle,13,'muted')
    rect(s,.6,7.04,12.13,.012,'line');tx(s,.6,7.13,11.6,.22,source,8.4,'muted');tx(s,12.25,7.08,.5,.28,f'{i:02}',11,'muted',align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text=notes+'\n\n证据来源：'+source
    SLIDES.append({'number':i,'title':title,'subtitle':subtitle,'source':source,'notes':notes,'figures':[]})
    return s

def pic(slide,path,x=.6,y=1.8,w=12.1,h=4.55):
    picture(slide,path,x,y,w,h);SLIDES[-1]['figures'].append('assets/'+Path(path).name)

def banner(slide,value,color='pale',y=6.39):
    rect(slide,.6,y,12.13,.48,color);tx(slide,.78,y+.085,11.8,.3,value,14,'ink',bold=True)

def card(slide,x,y,w,h,kicker,value,body,color='teal'):
    rect(slide,x,y,w,h);rect(slide,x,y,.045,h,color)
    tx(slide,x+.22,y+.2,w-.44,.4,kicker,14,'muted');tx(slide,x+.22,y+.76,w-.44,.66,value,31,color,True)
    tx(slide,x+.22,y+1.65,w-.44,h-1.8,body,16)

def table(slide,headers,data,widths,x=.6,y=1.8,rowh=.58,size=16):
    xx=x
    for h,w in zip(headers,widths):rect(slide,xx,y,w,rowh,'navy');tx(slide,xx+.12,y+.14,w-.24,rowh-.15,h,size,'white',True);xx+=w
    for j,row in enumerate(data):
        xx=x
        for v,w in zip(row,widths):
            rect(slide,xx,y+(j+1)*rowh,w,rowh,'light' if j%2==0 else 'white')
            tx(slide,xx+.12,y+(j+1)*rowh+.13,w-.24,rowh-.12,v,size);xx+=w

# 01
s=base('语义三维码垛：项目进展与关键卡点','NBU_Papper｜阶段0 → O3数值稳定V2｜2026年10月21日',
       '最新证据：EXP_OPT_O3_LAYOUT_FLASH_V2_VALIDATION_20261020_01 / acceptance_report.json',
       '本次汇报围绕三件事：工程目前做到哪里，现有实验能支持什么结论，以及下一步应优先解决什么。最新V2真实完整配对已通过工程验收，不能继续沿用早前“运行中”的状态。所有图表来自已有实验和日志，不新增模型调用。')
tx(s,.72,1.94,11.8,1.45,'在线闭环已完成工程验证\n布局改善已有开发证据，模型增量仍待证明',33,bold=True)
for x,k,v in [(0.75,'最新完整配对','240 episode'),(4.95,'真实 Flash 调用','1502 次'),(9.05,'整单无违规完成','7 → 12 / 20')]:
    tx(s,x,4.25,3.5,.45,k,16,'muted');tx(s,x,4.88,3.6,.7,v,29,'teal',True)
tx(s,.75,6.35,11.8,.38,'研究范围：真值状态、严格在线、受限政策；当前数据仍为已见开发来源。',15,'muted')

# 02
s=base('当前进度可以归纳为三句话','区分工程完成、实验观察与论文证据', '来源：O3真实验收、失败归因、V2完整配对指标',
       '第一，数据、候选、语言解析、执行、回放已闭环。第二，局部业务目标会导致后续布局失败，布局策略已改善整单表现。第三，Flash与强公开Parser完全一致，当前还不能声称模型贡献、视觉贡献或经验记忆贡献。')
card(s,.6,1.95,3.9,3.95,'已经完成','可复现实验闭环','数据 → 候选 → AST → 执行\n逐步证据、独立回放、验收\n最新真实矩阵工程 PASS')
card(s,4.72,1.95,3.9,3.95,'已经观察到','布局策略有效信号','V2无违规完成 7/20 → 12/20\n观察模式违规实例 29 → 17\n仍存在完成率与合规的取舍','blue')
card(s,8.84,1.95,3.9,3.95,'仍需解决','模型增量与泛化','Flash / Parser 1502步一致\n仅10个独立来源订单\nO4视觉、O5记忆尚待实现','amber')
banner(s,'汇报主线：工程能力 → 失败机制 → 改进结果 → 研究卡点。')

# 03
s=base('工程已推进到 O3 布局验证阶段','旧阶段1D≈在线闭环；旧1E≈经验记忆，协议与验收仍需分别理解',
       '来源：docs/README.md；阶段0、1A—1C、O2b、O3对应验收报告',
       '早期先建设数据与几何平台，1C加入固定题文本和视觉模型。O1b至O2b逐步转向政策AST与共同执行器。O3完成完整订单闭环后，工作重点已转为失败归因、布局取舍和可复现性。O4与O5属于后续路线。')
table(s,['阶段','已形成的能力或证据','当前边界'],[
 ['阶段0 / 1A—1B','数据、几何候选、在线基线、回放','属性含合成值；未复现完整GOPT'],
 ['1C','360题 × 6方法；Text / Visual','固定题；旧test已查看'],
 ['O1b / O2a / O2b','政策题、绑定协议、AST＋Executor','强Parser已达到同一候选内上界'],
 ['O3在线','五方法完整订单；安全 / 业务观察','模型轨迹与Parser一致'],
 ['O3布局 / V2','失败归因、消融、数值稳定、真实配对','已见开发来源；需独立验证'],
 ['O4 / O5 / O6','视觉 / 经验记忆 / 机器人路线','尚未形成对应完整实测闭环']], [2.0,5.45,4.68],rowh=.62,size=15)

# 04
s=base('当前架构：语言解析与数值选址分工明确','Flash负责受限AST；几何安全、载荷与排序由确定性程序负责',
       '源码：runners/layout_flash_validation_v2.py；o3_online_adapter.py；selectors/policy_executor_v2.py',
       '模型输入是公开请求，不含隐藏真值。Parser和Flash生成的AST进入相同执行器，因此两者差异才可归因于解析来源。状态由程序提交候选后更新，隐藏政策用于事后评价，不进入普通方法的选址。')
nodes=[(.7,2.05,'当前状态','托盘＋已放箱＋当前箱'),(4.65,2.05,'物理候选','极值点 → 过滤 → K≤16'),(8.6,2.05,'公开请求','自然语言＋JSON＋Cxx'),
       (8.6,4.15,'Parser / Flash','指令 → 受限政策 AST'),(4.65,4.15,'共同执行器','业务集合 → 软目标 → 布局'),(.7,4.15,'检查与提交','业务模式＋独立物理检查')]
for x,y,a,b in nodes:
    rect(s,x,y,3.65,1.22,'light');tx(s,x+.18,y+.18,3.3,.38,a,21,'teal',True);tx(s,x+.18,y+.75,3.3,.35,b,14)
for x,y in [(4.35,2.42),(8.3,2.42)]:tx(s,x,y,.35,.4,'→',23,'teal')
for x,y in [(4.35,4.52),(8.3,4.52)]:tx(s,x,y,.35,.4,'←',23,'teal')
tx(s,11.8,3.4,.7,.5,'↓',25,'teal');tx(s,.75,5.75,11.8,.4,'提交后得到真实新状态，再进入下一箱；逐步日志用于独立评价与回放。',18)
banner(s,'当前主线 buffer_size=1；模型不生成坐标、不读取未来箱体。')

# 05
s=base('数据已经齐备，真实性与可见性必须分开','140个场景文件 = 100基础＋20语言配对＋20缓存配对',
       '来源：semantic_pallet_dataset_v1/README.md、splits/、configs/mvd_v0.1.json',
       'BED-BPP提供尺寸、质量和顺序。材质、易碎、可堆叠、品类及承载上限包含合成值。缓存场景和连续流存在，不代表缓存决策和动态经验已实现。当前O3只使用当前箱体和已放状态。')
card(s,.6,1.9,3.3,3.95,'数据规模','360 固定题','100个基础订单场景\ntrain 50 / validation 10\ntest-ID 20 / OOD各10')
table(s,['字段 / 资源','来源与当前使用方式'],[
 ['尺寸、质量、到货顺序','BED-BPP原始订单'],['易碎、材质、品类、承载','固定种子合成；非真实抗压测量'],
 ['隐藏规则 / Oracle','评价或明确标记的特权基线使用'],['当前模型输入','公开指令、真值状态、当前候选'],
 ['buffer3 / 连续任务流','数据已存在；当前主线尚未使用']], [2.65,5.8],x=4.28,y=1.95,rowh=.62,size=15)
banner(s,'当前结果证明系统与受限任务上的行为，尚不能直接外推真实仓储效果。')

# 06
s=base('最新实验是严格配对的 240 episode','同一候选器、同一执行器；只比较解析来源与三个布局变体',
       '来源：V2完整配对 manifest.json、data/tasks.private.json、audit/source_split.json',
       '矩阵为10来源订单×2政策×3布局×2模式×2解析来源，共240个episode。每个方法每模式每变体有20个政策任务，但独立来源仍为10个。两套政策分别采用1kg低位优先和3kg居中优先。')
parts=[('10','来源订单'),('2','公开政策'),('3','布局变体'),('2','执行模式'),('2','解析来源')]
for i,(v,label) in enumerate(parts):
    x=.7+i*2.48;rect(s,x,2.05,2.12,1.6,'light');tx(s,x+.15,2.23,1.82,.65,v,38,'teal',True,PP_ALIGN.CENTER);tx(s,x+.1,3.03,1.92,.35,label,16,align=PP_ALIGN.CENTER)
    if i<4:tx(s,x+2.14,2.53,.3,.4,'×',21,'muted')
table(s,['政策','业务限载','软目标优先级'],[['P0','易碎 1 kg','先重箱低位，再重箱居中'],['P1','易碎 3 kg','先重箱居中，再重箱低位']], [2.0,3.3,6.83],y=4.05,rowh=.58,size=16)
banner(s,'统计簇是10个来源订单；240 episode或1502次调用不能当作独立样本量。')

# 07
s=base('原 O3 暴露了“完成”与“合规”的冲突','历史业务观察模式；每方法20任务，10个来源订单',
       '来源：EXP_OPT_O3B_BUSINESS_OBSERVE_20261016_REMOTE_01 / metrics/summary.json',
       'Geometry完成20/20但无违规完成6/20；Parser、Flash和真值AST均完成11/20、无违规完成7/20。较低违规率伴随较低完成率，不能只强调违规改善。三种语义方法完全一致，指向共同执行器的长期布局问题。')
pic(s,historical_chart,y=1.75,h=3.95)
tx(s,.85,5.82,11.65,.38,'最终违规实例：Geometry 44｜固定属性 33｜Parser / Flash / 真值 AST 各29',17)
banner(s,'工程 PASS ≠ 业务最优；单步规则正确 ≠ 整单能够完成。')

# 08
s=base('真实案例：同样放完17箱，仍需检查业务违规','历史 SCN_0054 / P0；五方法完整订单可视化',
       '来源：report/visualization/O3_完整订单_SCN0054_P0_26-10-17/final_layouts.png',
       '这张图复用已经完成的五方法案例。所有方法均放完17箱；Geometry最终3个违规实例，其余方法2个。Parser、Flash、真值AST布局相同。原图中的“累计业务违规”此处按对应最终状态的违规实例理解，不能等同违规步骤数。橙色易碎、蓝色重箱、青色电子品，红框标识参与违规的箱体。')
pic(s,ASSETS/'09_历史五方法完整订单.png',y=1.63,h=4.78)
banner(s,'案例用于解释业务评价；它属于历史原策略，不是最新布局V2的效果图。')

# 09
s=base('卡点一：居中偏好会挤碎可用空间','历史 SCN_0052 / P0，第2箱首次分叉；两个动作当时均无业务硬违规',
       '来源：O3_FAILURE_ATTRIBUTION_20261019_02 / divergence_details.json；历史逐步日志',
       '第2箱后，两方法地面净空总面积同为0.8099平方米，但Parser的最大连续矩形仅0.3159平方米，Geometry为0.7200。Parser更偏向居中，代价是未来箱体适配空间被切碎。后续第7箱到货时，Parser停止于6/8。该图直接从保存位姿绘制；虚线是诊断报告给出的最大矩形。')
pic(s,divergence_chart,y=1.73,h=3.9)
tx(s,.8,5.78,11.6,.45,'同样净空总面积 0.8099 m²，最大连续矩形相差约2.28倍；原Parser最终仅放6/8箱。',17)
banner(s,'总面积相同不代表形状可用；不能只优化空地面积或继续增加K。')

# 10
s=base('卡点二：空地更多，也可能因支撑布局失败','历史 SCN_0054 / P1：第6箱分叉，第12箱到货时无物理候选',
       '来源：O3_FAILURE_ATTRIBUTION_20261019_02 / figures/O3_SCN_0054_P1.png、terminal_gate_funnel.csv',
       '该案例是“只保留地面面积”的反例。Parser在首次分叉时保留更多地面净空，却更早堆高。终止时Parser/Geometry垛高670/420mm，同高可堆叠连通面119900/226200平方毫米。70个位姿通过碰撞后全部支撑不足。故需要同时考虑形状、同高支撑、承载和高度。')
pic(s,ASSETS/'08_历史0054失败归因.png',y=1.65,h=4.72)
banner(s,'失败前垛高 670 vs 420 mm；70个无碰撞位姿最终全部因支撑不足被拒。')

# 11
s=base('离线干预支持布局假设，但不能解释全部失败','只替换首次分叉动作，后续继续原Parser＋Executor；非新方法成绩',
       '来源：O3_FAILURE_ATTRIBUTION_20261019_02 / single_action_interventions.csv',
       '九次单动作干预中，六个任务多放箱，四个完整放完，其中两个完整无违规；三任务放置数量未变。因而首次分叉对部分失败具有解释力，但不是唯一原因。注意九个政策任务仅来自六个独立订单，这里是离线因果诊断，不是新的可部署策略。')
pic(s,intervention_chart,y=1.9,h=3.8)
tx(s,.8,5.82,11.7,.38,'6/9 多放箱｜4/9 完整放完｜2/9 无违规完成｜3/9 放置数量未变',19,'teal',True)
banner(s,'诊断方向：保留可适配形状、同高支撑与承载；局部居中不能代表长期布局质量。')

# 12
s=base('布局改进：先明确目标取舍，再比较空间代理','三变体使用同一物理候选与政策；预算控制臂分离软目标放松的收益',
       '源码：selectors/layout_preserving_v2.py；evaluation/layout_preservation.py',
       'original维持原软目标字典序。两个relaxed变体均允许相同软预算：heavy_low 0.05、heavy_center 0.25。relaxed_geometry仅按几何选择；relaxed_wo_floor在允许集合内按support、capacity、height均值排序。相对original的收益包含预算取舍，相对同预算几何控制的差异才对应布局代理。布局特征是启发式代理，不是未来可放性证书。')
card(s,.6,1.95,3.9,3.83,'original','原策略','业务可行优先\n两级软目标严格排序\n最后按几何与Cxx破同分','muted')
card(s,4.72,1.95,3.9,3.83,'relaxed_geometry','预算＋几何','与原策略比较软目标损失\n在允许集合内选几何最高\n控制“放松目标”的收益','blue')
card(s,8.84,1.95,3.9,3.83,'relaxed_wo_floor','预算＋布局','使用相同预算和允许集合\nsupport / capacity / height\n均值评分，移除地面项')
banner(s,'软预算：heavy_low=0.05，heavy_center=0.25；不读取未来箱体。')

# 13
s=base('数值问题已修复；旧失败证据继续保留','V1完整矩阵曾因跨Python排序差异失败，V2重新定义稳定排序',
       '来源：O3_LAYOUT_FLASH_BC_REVIEW_20261020_01；O3_LAYOUT_NUMERIC_V2_20261020_01',
       '同一三项特征用sum求均值，Python3.11得到0.4207292622421585，3.13得到0.42072926224215856。V1按原始浮点值排序，使近同分动作改变并引发整单分叉。V2使用fsum，先求集合最大值，再保留M减1e-12内候选，之后稳定比较几何分并按Cxx选择。该精度仅作用排序，不扩大业务或物理容差。')
rect(s,.6,1.95,5.7,3.83,'warm');tx(s,.85,2.17,5.2,.45,'V1：末位差异可以改变离散动作',21,'red',True)
tx(s,.85,3.0,5.2,1.5,'Python 3.11：0.4207292622421585\nPython 3.13：0.42072926224215856\n同一特征，直接按浮点值排序产生分叉。',18)
tx(s,.85,5.0,5.2,.5,'V1 C：240 episode；工程验收 FAIL',18,'red',True)
rect(s,6.55,1.95,6.18,3.83,'pale');tx(s,6.8,2.17,5.65,.45,'V2：明确并冻结稳定排序协议',21,'teal',True)
tx(s,6.8,3.0,5.65,1.8,'math.fsum 均值\n集合最大值锚定，近同分阈值1e-12\n几何同分层 → Cxx升序\n原控制、业务集合与软预算维持原定义',18)
banner(s,'双环境各120任务，1521条记录的动作与状态交叉一致；旧V1 FAIL未被改写。')

# 14
s=base('最新真实完整配对已通过工程验收','V2实验闭环：离线认证 → 真实冒烟 → 真实完整配对',
       '来源：V2完整配对 acceptance_report.json、replay_report.json、api_usage.json',
       '10月21日完整矩阵已经完成，因此本报告更新早前交接状态。240个episode全部记录，3042条事件回放通过；与离线稳定基线比较1521条Parser事件。1502次调用均首答合法，无修复、重试或失败。审计检查物理违规、输入泄漏、预算越界均为零。这里的零物理违规不代表零业务违规。')
for x,k,v,b in [(.6,'真实完整矩阵','240 / 240','工程验收 PASS'),(4.72,'事件回放','3042 条','全部通过；基线对照1521条'),(8.84,'Flash真实调用','1502 次','首答合法；修复/失败/重试为0')]:card(s,x,1.95,3.9,3.95,k,v,b)
banner(s,'物理违规=0｜输入泄漏=0｜软预算越界=0；仍属探索性开发验证。')

# 15
s=base('安全模式：无违规整单成功从7/20升至12/20','最新V2真实矩阵；Parser与Flash指标完全相同，此处只画一组',
       '来源：V2完整配对 metrics/summary.json；每行20任务、10个独立来源',
       '安全模式中所有提交均无业务违规。原策略7/20完整无违规，预算几何8/20，预算布局12/20。布局对原策略增加25个百分点，但其中包含预算取舍；对同预算几何控制增加20个百分点。平均件数比例上布局86.43%低于几何88.40%，说明仍有取舍。')
pic(s,safe_chart,y=1.85,h=3.9)
tx(s,.8,5.82,11.7,.42,'预算＋布局较原策略 +25个百分点；较同预算几何 +20个百分点。',19,'teal',True)
banner(s,'无违规提交不等于整单成功；仍有8/20任务未完成。')

# 16
s=base('业务观察模式：完成率提高，合规取舍仍存在','最新V2真实矩阵；相同预算下比较几何与布局策略',
       '来源：V2完整配对 metrics/summary.json、per_episode.csv',
       '预算几何和预算布局均完成18/20。布局无违规完成12/20，几何8/20；最终违规实例为17对30。但平均件数比例布局94.12%，几何99.41%。业务违规步骤率分别21/264=7.95%和29/282=10.28%，不能把最终违规实例数当作违规步骤数。')
pic(s,observe_chart,y=1.8,h=3.95)
tx(s,.8,5.82,11.7,.42,'预算＋布局：21/264违规步（7.95%）；平均件数94.12%，低于几何99.41%。',17)
banner(s,'不能称全面占优：无违规完成更好，但残留失败任务的放置比例更低。')

# 17
s=base('改善案例：SCN_0054 / P1 从11箱提升到17箱','最新V2业务观察模式；图像直接由保存的真实位姿重绘',
       '来源：V2完整配对 episodes/observe_business/*/public_parser/O3_SCN_0054_P1/steps.jsonl',
       '原策略11/17且无违规，但第12箱到货时无物理候选。预算几何完成17/17，但最终2个业务违规实例；预算布局完成17/17且最终0违规。安全模式下同一任务分别11/17、15/17、17/17。这表明布局代理在该订单上同时改善完成和合规，但单个案例不能证明泛化。颜色：橙色易碎，青色非易碎重箱，灰色其他；不标示违规身份。')
pic(s,case54,y=1.85,h=3.7)
tx(s,.82,5.53,11.5,.23,'颜色：橙=易碎；青=非易碎重箱；灰=其他。颜色不表示违规；三图采用相同视角和坐标尺度。',10,'muted')
tx(s,.8,5.87,11.7,.4,'原策略：未完成且零违规  →  预算＋布局：完整完成且零违规。',19,'teal',True)
banner(s,'同订单安全模式：11/17 → 15/17 → 17/17；布局提升可在真实轨迹中解释。')

# 18
s=base('残留卡点：SCN_0060 / P1 仍无法完成','最新V2业务观察模式；“更少违规”无法抵消提前停止',
       '来源：V2完整配对 episodes/observe_business/*/public_parser/O3_SCN_0060_P1/steps.jsonl',
       '预算几何放16/17但最终2个业务违规；预算布局只放7/17且零违规，两者最终都因无物理候选停止。P0同样为16/17对7/17，但几何最终3个违规。当前support/capacity/height代理不能充分表示长宽适配和后续可放性；这只是后续设计方向，不是已证实修复方案。单动作干预对0060/P0也没有增加放置数量。')
pic(s,case60,y=1.85,h=3.7)
tx(s,.82,5.53,11.5,.23,'颜色：橙=易碎；青=非易碎重箱；灰=其他。颜色不表示违规；图示为各策略最终状态。',10,'muted')
tx(s,.8,5.87,11.7,.4,'几何预算控制：16/17、2违规；布局策略：7/17、0违规；两者均未完成。',18)
banner(s,'下一步需研究带长宽、高度、支撑与承载条件的可放性代理，并保留失败分母。')

# 19
s=base('违规减少13个：8个来自相同完成订单的改善','最新V2观察模式；两个改进方案完成完全相同的18个任务，各提交250箱',
       '来源：V2完整配对 metrics/per_episode.csv；本报告按相同完成任务独立拆分',
       '不能把全部违规减少都归因于少放箱子。两个变体完成完全相同的18个任务，各提交250箱；这些任务上的最终违规从25降到17，减少8个。另5个减少来自失败的0060双政策，放置进度由32箱降至14箱。因此总减少13个中，8个来自相同完成订单的布局改善，5个与不同暴露有关。这是描述性拆分，不是预注册子集检验。相对原策略新增五个无违规任务来自三个来源。同预算无违规完成率差20个百分点的探索区间为[0,50]，仍需独立来源验证。')
pic(s,decomposition,y=1.95,h=3.65)
tx(s,.8,5.72,11.7,.5,'最终违规总数 30 → 17：相同完成任务减少8个；失败订单放置暴露差异减少5个。',18,'teal',True)
banner(s,'收益并非全由少放箱造成；子集拆分仅作解释，独立来源验证仍是关键。')

# 20
s=base('模型增量尚未出现：1502步完全一致','最新真实矩阵中，Flash与充分开发的Public Parser进入共同Executor',
       '来源：V2完整配对一致性CSV；O3_LAYOUT_FLASH_V2_RESULTS_20261021_01独立AST复核',
       '当前有效政策、同一Flash状态上的动作、闭环对应状态这三个指标均为1502/1502一致。原始AST结构并非全部相同：完整矩阵有52个条件AST，但其then/else分支均相同，约简后语义与Parser一致。这不构成真实条件推理证据。它证明Flash在受限政策文本上的接口可靠，却不能证明优于强规则解析器。该一致性仅对应当前固定P0/P1任务。后续应冻结Parser后构造未见表达与政策组合，并分别测AST、语义执行和整单质量。')
for x,k in [(.6,'当前有效政策一致'),(4.72,'同状态动作一致'),(8.84,'闭环状态一致')]:card(s,x,1.95,3.9,3.5,k,'1502 / 1502','当前P0/P1受限政策任务\n三变体、双执行模式')
tx(s,.8,5.78,11.7,.4,'原始AST并非全同：C含52个冗余条件包装，双分支相同，不构成条件推理证据。',16)
banner(s,'当前输入为Text＋JSON；3D展示不构成视觉决策贡献，也尚无动态经验记忆收益。')

# 21
s=base('当前卡点与优先级已经明确','数值复现卡点已解决；核心研究卡点仍在布局与模型贡献',
       '依据：当前实验事实；右列为建议交付，尚未实施',
       '优先级一是布局泛化和残留失败：冻结三策略，在独立来源验证并诊断0060。优先级二是模型增量：冻结强Parser，按语言模板与来源分组隔离。之后再落实视觉独有信息和经验改变动作的接口。历史O1b/O2a人工签认与O2b统计设计仍是未结项，不能改写为通过。')
table(s,['优先级 / 卡点','当前证据','下一项具体交付'],[
 ['P1 布局与整单取舍','0060仍失败；代理不全面','独立来源配对＋0060终止漏斗'],
 ['P1 泛化与统计','仅10个已见来源；消融筛选','冻结策略、主终点、来源与统计方案'],
 ['P2 模型独立贡献','1502步与Parser完全一致','未见语言/政策组合，冻结强Parser'],
 ['P3 视觉贡献','O3只用Text＋JSON','图像独有属性与OCR等公平对照'],
 ['P3 动态经验','尚无可改变动作的记忆接口','定义检索、策略影响、验证与失效'],
 ['后续 执行迁移','目前是确定性真值环境','再接仿真、感知与机器人验证']], [2.5,4.2,5.43],rowh=.62,size=14)

# 22
s=base('后续推进顺序：先固定证据，再扩展能力','建议路线；本报告不执行新实验、不改动算法',
       '依据：O3失败归因、V2真实结果与现有研究路线',
       '首先归档当前V2真实完成结果并统一导航状态；随后固定三策略和计算预算，在独立来源验证，保留0060等负例；并行准备语言泛化协议，但不看到测试失败后继续修Parser；最后根据真实信息缺口推进O4，再定义O5经验作用于决策的接口。正式test前冻结来源、模板、主指标、重复次数与统计方法。')
steps=[('01','整理当前证据','V2结果、版本、指标、案例一致\n保留V1 FAIL与既有负结果'),('02','验证布局泛化','独立来源、固定预算与变体\n终止漏斗＋完整无违规主指标'),('03','验证模型价值','未见表达与政策组合\n强Parser冻结后再测试'),('04','再扩展视觉与记忆','先找图像独有信息\n先定义经验如何改变动作')]
for i,(n,t,b) in enumerate(steps):
    y=1.86+i*1.08;rect(s,.7,y,.73,.73,'teal');tx(s,.81,y+.15,.5,.4,n,22,'white',True)
    tx(s,1.7,y+.02,3.4,.55,t,23,bold=True);tx(s,5.1,y+.02,7.4,.77,b,17)
banner(s,'正式新test前冻结协议；以来源订单为统计簇，继续报告未完成与API失败。')

# 23
s=base('汇报结论：工程可用，算法有信号，贡献待分离','建议对外表述的结论与证据边界',
       '来源：本次汇报全部原始验收与指标；不将工程PASS等同论文结论',
       '可以报告：完整在线框架和V2真实配对通过工程验收；局部政策与长期布局存在取舍；已见开发来源上预算布局改善无违规整单完成。不能报告：Flash优于强Parser、已经验证视觉或记忆收益、已经得到独立test结论、已验证机器人执行。最值得投入的下一步是独立来源的布局验证与模型语言泛化实验。')
rect(s,.6,1.95,5.88,3.95,'pale');tx(s,.85,2.2,5.4,.45,'已有证据支持',24,'teal',True)
tx(s,.85,3.0,5.35,2.45,'完整订单在线实验闭环\nV2真实完整配对工程 PASS\n开发集无违规完成 7/20 → 12/20\n长期布局取舍有可追溯案例',21)
rect(s,6.74,1.95,5.99,3.95,'warm');tx(s,6.99,2.2,5.5,.45,'仍需补齐证据',24,'red',True)
tx(s,6.99,3.0,5.45,2.45,'模型相对强Parser的增量\n独立来源与正式新test\n视觉、经验记忆的独立贡献\n仿真与机器人执行有效性',21)
banner(s,'下一阶段主问题：当前布局收益是否泛化，模型在哪类信息或语言缺口上真正有用？')

# 24 appendix
s=base('附录：指标、版本与证据入口','完整讲解稿、逐项数据和可点击路径见配套Markdown',
       '证据根目录：semantic_pallet_planner/outputs/；诊断附件：report/analysis/',
       '完整放完指放置箱数达到订单总数；无违规完成要求同时完整放完且最终业务违规为零；违规步是新增或加重违规的提交步骤，实例数则是最终状态规则对象计数，两者不能互换。mean item ratio为任务放置比例的平均。V1和V2布局协议不可混用。本次仅重绘已有日志，不调用API。原始来源的SHA256记录在evidence.json，构建脚本可在本目录复现图表与PPT。')
table(s,['项目','定义 / 入口'],[
 ['完成 / 无违规完成','放完全部箱体 / 放完且最终业务违规为0'],['违规步 / 最终违规实例','新增或加重违规的步骤 / 最终状态违规对象计数'],
 ['最新真实实验','EXP_OPT_O3_LAYOUT_FLASH_V2_VALIDATION_20261020_01'],['当前排序协议','layout_preserving_v2_fsum_max_anchor_1e-12'],
 ['失败诊断','report/analysis/O3_FAILURE_ATTRIBUTION_20261019_02'],['历史数值失败 / V2认证','O3_LAYOUT_FLASH_BC_REVIEW / O3_LAYOUT_NUMERIC_V2']], [3.0,9.13],rowh=.62,size=14)

PPTX=HERE/'NBU_Papper项目进展与卡点汇报_20261021.pptx'
prs.save(PPTX)

def render_preview():
    """Approximate rasterization for QA, not an Office rendering engine."""
    out=HERE/'preview';out.mkdir(exist_ok=True)
    W,H=1600,900;scale=W/prs.slide_width
    fontcache={}
    def ff(size,bold=False):
        key=(size,bold)
        if key not in fontcache:fontcache[key]=ImageFont.truetype(BOLD if bold else FONT,size)
        return fontcache[key]
    def col(color,default='#172D43'):
        try:return '#'+str(color.rgb)
        except Exception:return default
    issues=[];thumbs=[]
    for i,slide in enumerate(prs.slides,1):
        im=Image.new('RGB',(W,H),'white');d=ImageDraw.Draw(im)
        for sh in slide.shapes:
            x,y,w,h=[int(v*scale) for v in (sh.left,sh.top,sh.width,sh.height)]
            if sh.shape_type==MSO_SHAPE_TYPE.PICTURE:
                p=Image.open(BytesIO(sh.image.blob)).convert('RGB').resize((max(1,w),max(1,h)),Image.Resampling.LANCZOS);im.paste(p,(x,y))
            elif sh.shape_type==MSO_SHAPE_TYPE.AUTO_SHAPE:
                try:
                    if sh.fill.type:d.rectangle((x,y,x+w,y+h),fill=col(sh.fill.fore_color,'white'))
                except Exception:pass
            if sh.has_text_frame and sh.text:
                tf=sh.text_frame;yy=y
                for para in tf.paragraphs:
                    if not para.runs:continue
                    r=para.runs[0];size=int((r.font.size.pt if r.font.size else 18)*W/960)
                    f=ff(size,bool(r.font.bold));lines=[];cur=''
                    for ch in para.text:
                        if cur and d.textlength(cur+ch,font=f)>w:
                            lines.append(cur);cur=ch
                        else:cur+=ch
                    if cur:lines.append(cur)
                    for l in lines:
                        xx=x
                        if para.alignment==PP_ALIGN.CENTER:xx=x+(w-d.textlength(l,font=f))/2
                        if para.alignment==PP_ALIGN.RIGHT:xx=x+w-d.textlength(l,font=f)
                        d.text((xx,yy),l,font=f,fill=col(r.font.color),anchor='lt');yy+=size*1.13
                    yy+=5*W/960
                if yy-y>h+12:issues.append({'slide':i,'text':sh.text[:70],'estimated_px':yy-y,'box_px':h})
        im.save(out/f'slide_{i:02}.png')
        thumbs.append(im.resize((480,270),Image.Resampling.LANCZOS))
    sheet=Image.new('RGB',(4*496,math.ceil(len(thumbs)/4)*300),(226,233,238));d=ImageDraw.Draw(sheet)
    for i,im in enumerate(thumbs):
        x=(i%4)*496+8;y=(i//4)*300+8;sheet.paste(im,(x,y));d.text((x,y+273),f'{i+1:02d}',font=ff(13),fill='#607487')
    sheet.save(HERE/'PPT预览.png')
    return issues

preview_issues=render_preview()
print(json.dumps({'pptx':str(PPTX),'slides':len(prs.slides),'preview_estimated_overflows':preview_issues},ensure_ascii=False,indent=2))

# Report narrative is intentionally richer than the slide text.
MD=HERE/'NBU_Papper项目进展与卡点报告_20261021.md'
intro='''# NBU_Papper 项目进展与卡点报告

汇报日期：2026-10-21。面向组会或阶段汇报；建议主讲20—30分钟，最后一页为指标附录。

本报告依据仓库源码、已有原始实验、验收文件与逐步轨迹编写。此次不运行新实验，不调用API，不改动算法与历史产物；图表为既有实验数据重绘，案例来自实际保存位姿。最新状态已更新为：**数值稳定V2真实完整配对240个episode工程验收PASS**。

- [下载PPTX](NBU_Papper项目进展与卡点汇报_20261021.pptx)
- [24页PPT总览](PPT预览.png)
- [来源摘要与校验记录](evidence.json)

## 1. 面向汇报的核心结论

项目已经从固定候选的单步语义选择，推进到完整订单的属性感知在线闭环。数据、几何候选、语言政策解析、确定性执行、独立物理检查、业务评价和日志回放已具备工程证据。

当前主要工作已转为：解释局部政策排序为何导致长期布局失败，并在明确的软目标预算内改进整单表现。V2去地面项布局策略在已见开发数据上把无违规整单完成从7/20提高到12/20；但相对同预算几何控制仍有放置比例取舍，SCN_0060等失败尚未解决。

模型的独有贡献尚未出现。最新1502条比较记录中，Flash与Public Parser的当前有效政策、同状态动作和闭环状态完全一致。当前结果支持受限政策解析接口可靠，不能支持Flash优于强Parser、视觉贡献或动态记忆贡献。原始AST结构并非全同：完整配对中有52个冗余条件AST，其双分支均相同，约简后与Parser语义一致；不能将其解释为真实条件推理。

## 2. 工程框架与数据交互

```text
BED-BPP订单＋合成属性
  → DatasetRepository只读加载与Schema/引用校验
  → 当前状态（托盘、已放箱、当前箱；buffer_size=1）
  → 极值点/支撑层枚举与物理过滤 → 多样性Top-K≤16
  → 公开请求（自然语言、当前事实、候选Cxx；不含隐藏答案）
  → Public Parser / Flash → 受限政策AST
  → 共同Executor / V2布局选择器
  → 业务模式检查＋独立物理检查 → 提交候选位姿
  → 新真实状态 → 下一箱
  → JSONL轨迹、CSV指标、manifest、独立回放和验收
```

| 代码位置 | 职责 |
|---|---|
| `repository.py`、`domain.py`、`config.py` | 只读数据与配置、不可变快照、状态恢复 |
| `planner/generator.py`、`constraints.py`、`scoring.py` | 候选枚举、硬约束与几何评分 |
| `selectors/instruction_parser_o1b.py` | 公开自然语言政策解析 |
| `selectors/policy_executor_v2.py` | 冻结的原政策数值执行器 |
| `selectors/layout_preserving_v2.py` | 相同预算允许集合上的稳定布局排序 |
| `vlm/policy_protocol_o2b.py`、`layout_flash_client.py` | AST协议、绑定校验、真实调用、最多一次格式修复 |
| `runners/attribute_o3_online.py` | 历史O3五方法在线矩阵 |
| `runners/layout_flash_validation_v2.py` | 最新三变体双解析来源矩阵与回放验收 |
| `evaluation/` | 独立物理检查、传播载荷、业务评价、布局代理和回放 |
| `logging/`、`visualization/` | 原始证据持久化与3D展示 |

早期1A—1C主要通过`cli.py`和`PalletEnvironment`运行，当前O3有独立模块入口；不能只读通用CLI推断最新执行逻辑。第三方GOPT、ReMe、PackLab目录属于参考资源，不代表完整算法已集成。

物理层限制0°/90°水平旋转，检查边界、总载荷、碰撞、支撑率（至少0.70）、重心投影、不可堆叠与递归承重。几何评分权重为紧凑0.35、高度0.25、稳定0.25、平衡0.15。独立物理复核不调用候选生成器。

Flash只返回AST，需绑定`request_id`、`state_hash`、`policy_digest`；不生成坐标，不读取未来箱体。当前`heavy_center`是各重箱中心距的质量加权平均，不能误称整垛质心偏移。

## 3. 数据与阶段进度

数据集包含100个基础场景（train50、validation10、test-ID20、Geometry-OOD10、Semantic-OOD10）、20个Language-OOD配对场景和20个buffer3配对场景，共140个场景文件；另有360个固定决策题、360组固定候选、360份Oracle、100份参考方案以及8条连续流。

尺寸、质量和到货顺序来自BED-BPP；材质、易碎性、可堆叠性、品类与承载上限包含固定种子合成值。合成承载上限与政策的业务限载是两种约束。buffer3和任务流已存在，但当前主线仍是buffer1，尚无动态经验记忆。

| 阶段 | 已完成 | 当前边界 |
|---|---|---|
| 0 / 1A—1B | 数据、接口、几何候选、在线基线、回放和可视化 | 非完整GOPT学习型基线 |
| 1C | 360题×6方法，2160条历史固定题记录；Text/Visual | 旧test已查看，视觉增量很小 |
| O1b / O2a | 语言政策题、AST、绑定校验与协议修复 | 历史人工签认等仍按原记录待完成 |
| O2b | 180题×6方法，1080条记录回放通过 | 统计设计待补；未证实模型超过强Parser |
| O3在线 | 安全和业务观察模式真实五方法工程PASS | 单步局部目标导致整单布局问题 |
| O3失败归因 | 9失败任务、首分叉、终止漏斗、单动作干预 | 不是连续空间全局无解证明 |
| O3布局 | 9变体×双政策×双模式，共360离线episode | 已见开发来源，较好变体经消融筛选 |
| O3数值V2 | 双Python认证、真实冒烟及240真实配对PASS | 仍无独立新test和模型增量证据 |
| O4 / O5 / O6 | 视觉 / 经验记忆 / 机器人研究路线 | 尚未形成对应完整实测闭环 |

## 4. 最新真实实验与结果

矩阵：**10个来源订单×2政策×3变体×2模式×2解析来源＝240 episode**。每行20个政策任务来自10个来源，不能把240 episode、1502调用或3042事件当成独立样本量。

P0：易碎业务限载1kg，重箱低位优先、居中次之。P1：限载3kg，居中优先、低位次之。两者包含相应禁止接触规则。`safe`在无业务可行候选时停止；`observe_business`允许提交业务违规并在无解时按最小伤害等排序。两模式均维持物理检查。

三变体：`original`保持原字典序执行器；`relaxed_geometry`在显式预算内选几何最高；`relaxed_wo_floor`使用相同预算，增加support/capacity/height三项均值布局评分。预算为heavy_low=0.05、heavy_center=0.25。相对原策略的变化包含目标放松，不能全部归因于布局评分。

| 工程项 | 最新结果 |
|---|---|
| 完整配对 | 240/240 episode，工程PASS |
| 回放 | 3042条事件PASS；对照稳定基线1521条 |
| 真实调用 | 1502次；首答合法1502/1502 |
| 修复 / 失败 / 传输重试 | 0 / 0 / 0 |
| 物理违规 / 输入泄漏 / 预算越界 | 0 / 0 / 0 |
| 当前有效政策 / 同状态动作 / 闭环状态一致 | 各1502/1502；不等于原始AST结构全同 |

Parser与Flash各组指标完全相同，下表合并列示；不能把相同结果计为额外独立证据。

| 模式 | 变体 | 完整放完 | 无违规完成 | 平均件数比例 | 违规步 / 提交步 | 最终违规实例 |
|---|---|---:|---:|---:|---:|---:|
'''
for mode in ['safe','observe_business']:
    for v in VARIANTS:
        d=summary[f'{mode}/{v}/public_parser']
        intro+=f'| {mode} | {v} | {d["completed"]}/20 | {d["clean_completed"]}/20 | {d["mean_placed_item_ratio"]:.2%} | {d["violating_steps"]}/{d["placed_steps"]} | {d["final_business_violation_instances"]} |\n'
intro+='''
![安全模式](assets/02_安全模式.png)

![业务观察模式](assets/03_业务观察模式.png)

解读：预算布局相对原策略无违规完成提升25个百分点，相对同预算几何控制提升20个百分点。观察模式最终违规17对30，但平均件数94.12%低于几何99.41%，不能称全面占优。聚类bootstrap属于探索性开发统计，未作多重比较修正。

同预算无违规完成率差为+20个百分点，探索性95%来源聚类区间为[0,50]个百分点，不能据此宣称稳健显著优势。

### 4.1 违规减少是否仅仅因为少放箱子

两个改进方案完成完全相同的18个观察模式任务，这批任务均提交250箱。按同一完成任务集合独立拆分：

| 相同任务集合 | 预算＋几何 | 预算＋布局 |
|---|---:|---:|
| 完成任务数 / 提交件数 | 18 / 250 | 18 / 250 |
| 这些完成任务最终违规实例 | 25 | 17 |
| 这些完成任务违规步骤 | 25 | 21 |
| 失败0060双政策提交件数 | 32 | 14 |
| 失败0060双政策最终违规实例 | 5 | 0 |

总违规实例减少13个，其中8个来自相同完成订单上的改善，5个来自失败订单的不同放置暴露。因此收益不全由少放箱造成，但整体违规率仍受放置数量影响。这是描述性解释，没有新增子集显著性检验。

![违规改善来源拆分](assets/10_违规改善来源拆分.png)

## 5. 重点案例及建议讲解顺序

### 5.1 历史五方法完整订单：完成不代表合规

SCN_0054/P0历史五方法均放完17箱；Geometry最终3个业务违规实例，其余方法2个，Parser/Flash/真值AST布局相同。此图用于建立“业务违规仍需单独检查”的直观认识，不能与最新V2布局效果混用。

![历史五方法](assets/09_历史五方法完整订单.png)

### 5.2 SCN_0052/P0：首次分叉如何切碎空间

历史原策略第2箱与Geometry分叉。两个备选动作当时均无业务硬违规，净空总面积同为0.8099m²，但Parser最大连续矩形0.3159m²、Geometry为0.7200m²。Parser最终停止于6/8箱。形状适配比总空地面积更有解释力。红虚线是诊断报告给出的最大地面连续矩形；视图来自历史保存位姿。

![首次分叉](assets/07_0052首次分叉.png)

### 5.3 SCN_0054/P1：支撑与堆高的反例

历史原策略第6箱分叉，首次分叉后保留更多地面净空，却更早堆高。第12箱到货前，Parser/Geometry垛高670/420mm；最大同高可堆叠连通面119900/226200mm²。70个无碰撞位姿全部因支撑不足被拒。不能仅靠地面空地面积解释失败。

9次单动作替换后，6个任务多放箱，4个完整完成，其中2个无违规完成；3个放置数量不变。这是离线机制干预，不是新策略成绩。

![干预](assets/04_单动作干预.png)

### 5.4 最新V2改善案例：SCN_0054/P1

| 模式 | 原策略 | 预算＋几何 | 预算＋布局 |
|---|---|---|---|
| safe | 11/17，最终0违规 | 15/17，最终0违规 | 17/17，最终0违规 |
| observe_business | 11/17，最终0违规 | 17/17，最终2违规 | 17/17，最终0违规 |

该案例展示从“未完成且零违规”到“完成且零违规”的提升。下图直接重绘最新真实矩阵的Parser轨迹；Flash指标和状态一致。橙色为易碎箱，青色为非易碎重箱，灰色为其他箱体；颜色不表示违规。

![0054最新布局](assets/06_SCN_0054_P1_真实布局.png)

### 5.5 最新V2残留失败：SCN_0060/P1

观察模式下，原策略7/17且最终0违规；预算几何16/17且最终2违规；预算布局7/17且最终0违规，均以无物理候选终止。P0同样是几何16/17、布局7/17，几何最终3违规。当前布局代理还不能充分表示长宽适配和长期可放性，零违规不能抵消未完成。

![0060残留失败](assets/06_SCN_0060_P1_真实布局.png)

### 5.6 收益集中性与统计边界

预算布局较原策略新增5个无违规完成任务：0051/P0、0051/P1、0054/P0、0054/P1、0057/P0，实际来自3个来源订单，原有成功任务未丢失。去地面项来自既有消融筛选，必须在独立来源复核。

相对同预算几何控制，新增的4个无违规完成任务集中在0051、0054两个来源，不能把政策任务数作为独立样本数。

![逐任务结果](assets/05_逐任务成功分布.png)

## 6. 数值复现问题与版本边界

V1真实冒烟72 episode、508调用工程PASS；V1完整配对240 episode、1495调用，因5条历史布局轨迹不一致保留FAIL。同一特征均值在Python3.11与3.13分别为0.4207292622421585和0.42072926224215856，微小数值差异改变近同分动作，进而造成整单分叉。

V2使用`math.fsum`和集合最大值锚定的1e-12近同分层，之后比较几何评分并按Cxx选择。排序协议为`layout_preserving_v2_fsum_max_anchor_1e-12`。它没有修改物理判定、业务EPS、原控制或软预算。

两环境各自从空托盘运行120个逻辑任务，1521条事件动作与状态交叉一致；这是同一矩阵的双环境复现，不是240个独立研究样本。最新V2真实冒烟72 episode、508调用已PASS，完整矩阵240 episode、1502调用也已PASS。

V2观察模式去地面项最终违规实例为17，不得引用旧布局15或V1 C的16来描述V2。V1的FAIL继续保留。

## 7. 当前卡点与后续建议

| 优先级 | 卡点 | 建议交付与判定重点 |
|---|---|---|
| P1 | 长期布局代理与目标取舍 | 独立来源比较三变体；0060终止漏斗；完整无违规、完成率、件数比例一起报告 |
| P1 | 泛化与统计 | 冻结来源、模板族、主终点、重复次数与聚类统计；旧已查看test仅作回归 |
| P2 | 模型相对强Parser的增量 | Parser充分开发后冻结；新语言、条件与政策组合；分别报告AST、执行语义、动作和整单效果 |
| P3 | 视觉贡献 | 先定义图像独有信息；比较部分信息Text、Visual、OCR＋同Executor与完整真值Text |
| P3 | 动态经验 | 先定义经验如何改变排序或策略；再比较NoMemory、冻结检索和动态记忆，隔离未来反馈 |
| 后续 | 机器人迁移 | 确定性几何通过后再验证仿真、可达性、感知与执行；目前不作机器人成功声明 |

历史O1b/O2a人工签认、维护工时及O2b完整统计设计仍按原报告保留未结，不将Agent核查改写成人工签认，也不增设其为本轮O3工程门槛。当前数值卡点已解决，不应再次列为待修复；正式新test、视觉与记忆尚未完成。

## 8. 24页PPT配套讲解稿

以下逐页给出建议讲解重点与证据位置。PPT中也嵌入了相同讲者备注。

'''
for item in SLIDES:
    intro+=f'### 第{item["number"]:02d}页：{item["title"]}\n\n{item["subtitle"]}\n\n**讲解：** {item["notes"]}\n\n**来源：** {item["source"]}\n\n'
intro+='''## 9. 可核对的原始证据入口

'''
links=[(EXP/'acceptance_report.json','最新完整配对验收'),(EXP/'replay_report.json','最新回放'),(EXP/'api_usage.json','最新API记录'),
 (EXP/'metrics/summary.json','最新方法汇总'),(EXP/'metrics/per_episode.csv','最新逐单指标'),(EXP/'metrics/parser_flash_agreement.csv','解析与动作一致性'),
 (EXP/'metrics/paired_cluster_comparisons.json','探索性来源聚类比较'),(DIAG/'README.md','历史失败归因'),
 (ROOT/'report/analysis/O3_LAYOUT_FLASH_BC_REVIEW_20261020_01/README.md','历史V1 B/C复核'),
 (ROOT/'report/analysis/O3_LAYOUT_NUMERIC_V2_20261020_01/README.md','V2离线认证'),
 (ROOT/'docs/03_阶段产出分析/O3/O3数值稳定V2真实冒烟与配对验证结果分析_20261021.md','V2真实结果专项分析'),
 (ROOT/'report/analysis/O3_LAYOUT_FLASH_V2_RESULTS_20261021_01/README.md','V2真实结果独立复核附件'),
 (ROOT/'report/visualization/O3_完整订单_SCN0054_P0_26-10-17/index.html','历史五方法交互3D案例')]
import os
for p,label in links:intro+=f'- [{label}]({os.path.relpath(p,HERE)})\n'
intro+='''
## 10. 交付与复现说明

- PPT为16:9、24页，正文、结构图和表格可编辑；统计图和3D图以高分辨率PNG嵌入，独立PNG保存在`assets/`。
- 本次可视化仅重绘已有实验；没有模型生成图片，没有新增模型调用，没有重跑实验。
- `preview/`逐页PNG和`PPT预览.png`是本地近似排版预览，并非PowerPoint/LibreOffice原生渲染；实际播放以PowerPoint/WPS为准。
- `evidence.json`记录读取来源的SHA256、验收摘要、数据一致性与版面检查；`build_report.py`可在本机已有环境重建报告。
- 构建命令：`MPLCONFIGDIR=/tmp/nbu-report-mpl /home/cyw/anaconda3/envs/hilserl_sim/bin/python report/项目进展汇报_20261021/build_report.py`。
- 原始完整实验仍在`outputs/`，通常不随Git提交；单独复制Markdown时应同时复制`assets/`，原始证据链接在仓库内有效。
'''
MD.write_text(intro,encoding='utf-8')

case_evidence=[]
for sid in ['SCN_0054','SCN_0060']:
    for mode in ['safe','observe_business']:
        for v in VARIANTS:
            r=episode_row(mode,v,sid)
            case_evidence.append({k:r[k] for k in ['scenario_id','policy_id','mode','variant','placed_items','planned_items','final_business_violation_count','termination_reason']})

# Check source integrity after reading; all mutations are report-only.
source_changes=[name for name,h in SOURCES.items() if hashlib.sha256((ROOT/name).read_bytes()).hexdigest()!=h]
assert not source_changes
evidence={'generated_at':datetime.now(timezone(timedelta(hours=8))).isoformat(),
 'scope':'Existing experimental evidence only; no experiment rerun, no API calls, no algorithm changes',
 'latest_experiment':str(EXP.relative_to(ROOT)),'acceptance':acceptance,'replay':replay,'api':usage,
 'agreement_records':len(agree),'agreement_mismatches':{k:sum(r[k]!='True' for r in agree) for k in ['ast_equal','same_state_action_equal','closed_loop_states_equal']},
 'source_sha256':SOURCES,'source_changes_during_build':source_changes,'case_rows':case_evidence,
 'independent_aggregate_assertions':aggregate_checks,
 'slides':len(prs.slides),'notes_slides':len(SLIDES),'preview_renderer':'local approximate PIL renderer, not Office',
 'estimated_text_overflows':preview_issues,'api_calls_in_this_build':0}
(HERE/'evidence.json').write_text(json.dumps(evidence,ensure_ascii=False,indent=2)+'\n')
(HERE/'逐页提纲.json').write_text(json.dumps(SLIDES,ensure_ascii=False,indent=2)+'\n')

# Check the complete local slide geometry and the OOXML package.
for i,slide in enumerate(prs.slides,1):
    for sh in slide.shapes:
        assert sh.left>=0 and sh.top>=0 and sh.left+sh.width<=prs.slide_width+Inches(.01) and sh.top+sh.height<=prs.slide_height+Inches(.01),(i,sh.name)
with zipfile.ZipFile(PPTX) as z:
    assert z.testzip() is None
    assert sum(n.startswith('ppt/slides/slide') and n.endswith('.xml') for n in z.namelist())==24
assert len(Presentation(PPTX).slides)==24
import re
md_targets=re.findall(r'!?\[[^\]]*\]\(([^)]+)\)',intro)
missing_links=[t for t in md_targets if not (HERE/t).exists()]
assert not missing_links,missing_links
evidence.update(markdown_links_checked=len(md_targets),missing_markdown_links=missing_links,
                pptx_zip_integrity='PASS',slide_geometry_bounds='PASS',pptx_reopened_slides=24)
(HERE/'evidence.json').write_text(json.dumps(evidence,ensure_ascii=False,indent=2)+'\n')
print('Wrote',MD)
