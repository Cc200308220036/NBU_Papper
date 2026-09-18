#!/usr/bin/env python3
"""生成阶段1A—1B十页中文汇报PPT。"""

from pathlib import Path
import json

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
DATASET = ROOT / "semantic_pallet_dataset_v1"
PLANNER = ROOT / "semantic_pallet_planner"
EXP_1A = PLANNER / "outputs/EXP_1A_20260912_FINAL"
EXP_1B = PLANNER / "outputs/EXP_1B_20260912_FINAL"
OUTPUT = ROOT / "docs/阶段1A-1B阶段汇报.pptx"

NAVY = "102A43"
NAVY_2 = "183B5B"
BLUE = "2F6BFF"
TEAL = "16A89A"
AMBER = "F2A93B"
RED = "E35D6A"
INK = "172B4D"
MUTED = "52677D"
LIGHT = "F4F7FB"
LINE = "D8E2EC"
WHITE = "FFFFFF"
PALE_BLUE = "EAF1FF"
PALE_TEAL = "E7F7F4"
PALE_AMBER = "FFF4DE"
PALE_RED = "FDECEE"
FONT = "Noto Sans CJK SC"
MONO = "DejaVu Sans Mono"


def rgb(hex_value):
    return RGBColor.from_string(hex_value)


def add_rect(slide, x, y, w, h, fill, line=None, radius=True):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line or fill)
    if radius:
        try:
            shape.adjustments[0] = 0.08
        except Exception:
            pass
    return shape


def add_text(slide, x, y, w, h, text, size=18, color=INK, bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, font=FONT,
             margin=0.04, line_spacing=1.05):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.line_spacing = line_spacing
    for run in p.runs:
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
    return box


def add_lines(slide, x, y, w, h, lines, size=15, color=INK, bullet=False,
              gap=4, font=FONT, margin=0.05):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    for idx, item in enumerate(lines):
        if isinstance(item, tuple):
            text, item_color, item_bold = item
        else:
            text, item_color, item_bold = item, color, False
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = ("• " if bullet else "") + text
        p.space_after = Pt(gap)
        p.line_spacing = 1.1
        for run in p.runs:
            run.font.name = font
            run.font.size = Pt(size)
            run.font.bold = item_bold
            run.font.color.rgb = rgb(item_color)
    return box


def add_title(slide, index, title, subtitle=None):
    add_text(slide, 0.65, 0.36, 0.55, 0.42, f"{index:02d}", 15, BLUE, True,
             valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, 1.18, 0.27, 11.35, 0.58, title, 26, NAVY, True,
             valign=MSO_ANCHOR.MIDDLE)
    if subtitle:
        add_text(slide, 1.2, 0.82, 11.0, 0.34, subtitle, 10.5, MUTED)
    add_rect(slide, 0.65, 1.12, 12.0, 0.025, BLUE, radius=False)


def add_footer(slide, page):
    add_text(slide, 0.68, 7.18, 8.0, 0.18,
             "Semantic Pallet Planner · 数据集V0.1 · 阶段1A—1B", 8.5, MUTED)
    add_text(slide, 12.0, 7.13, 0.65, 0.23, f"{page}/10", 9, MUTED,
             True, PP_ALIGN.RIGHT)


def add_metric(slide, x, y, w, value, label, color=BLUE):
    add_rect(slide, x, y, w, 1.08, WHITE, LINE)
    add_rect(slide, x, y, 0.07, 1.08, color, radius=False)
    add_text(slide, x + 0.2, y + 0.12, w - 0.3, 0.43, value, 24, color, True)
    add_text(slide, x + 0.2, y + 0.59, w - 0.3, 0.27, label, 10.5, MUTED)


def add_picture_crop(slide, path, x, y, w, h, line=LINE):
    path = Path(path)
    with Image.open(path) as im:
        iw, ih = im.size
    target = w / h
    source = iw / ih
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    if source > target:
        visible = target / source
        pic.crop_left = pic.crop_right = (1 - visible) / 2
    elif source < target:
        visible = source / target
        pic.crop_top = pic.crop_bottom = (1 - visible) / 2
    border = add_rect(slide, x, y, w, h, fill=WHITE, line=line, radius=False)
    # 必须写成真正的<a:noFill/>。部分WPS/PowerPoint版本会忽略
    # python-pptx对solidFill透明度的设置，导致白色边框矩形遮住图片。
    border.fill.background()
    return pic


def add_flow_box(slide, x, y, w, h, title, detail, color, pale):
    add_rect(slide, x, y, w, h, pale, color)
    add_text(slide, x + 0.15, y + 0.13, w - 0.3, 0.33, title, 15, color, True,
             align=PP_ALIGN.CENTER)
    add_text(slide, x + 0.13, y + 0.53, w - 0.26, h - 0.64, detail, 10.2,
             MUTED, align=PP_ALIGN.CENTER)


def add_arrow(slide, x, y, w=0.4, h=0.35, color=BLUE):
    shape = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = rgb(color)
    shape.line.color.rgb = rgb(color)
    return shape


def base_slide(prs, index, title, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = rgb(LIGHT)
    add_title(slide, index, title, subtitle)
    add_footer(slide, index)
    return slide


def build():
    summary = json.loads((DATASET / "statistics/dataset_summary.json").read_text())
    stage0 = json.loads((DATASET / "statistics/stage0_acceptance_report.json").read_text())
    topk = json.loads((EXP_1B / "metrics/topk_selection.json").read_text())
    replay = json.loads((EXP_1B / "replay_report.json").read_text())

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1. 封面
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = rgb(NAVY)
    add_rect(slide, 0, 0, 0.16, 7.5, TEAL, radius=False)
    add_text(slide, 0.78, 0.72, 7.3, 0.35, "硕士课题阶段汇报 · 2026", 13, "8FD9D1", True)
    add_text(slide, 0.75, 1.25, 8.5, 1.55,
             "面向语义约束的\n三维机械臂码垛规划", 31, WHITE, True,
             valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, 0.8, 2.92, 8.2, 0.5,
             "数据集架构 · 阶段1A统一平台 · 阶段1B几何规划与可视化", 16, "C9D8E8")
    add_rect(slide, 8.75, 0.78, 3.85, 5.6, NAVY_2, "315A78")
    add_text(slide, 9.12, 1.1, 3.1, 0.42, "当前完成状态", 18, WHITE, True)
    stages = [("阶段0", "数据与接口验收", TEAL), ("阶段1A", "统一实验平台", BLUE),
              ("阶段1B", "几何基线与动画", AMBER)]
    for i, (name, detail, color) in enumerate(stages):
        yy = 1.78 + i * 1.18
        add_rect(slide, 9.1, yy, 0.68, 0.68, color, color)
        add_text(slide, 9.1, yy + 0.09, 0.68, 0.35, "✓", 22, WHITE, True,
                 PP_ALIGN.CENTER)
        add_text(slide, 9.98, yy - 0.01, 1.9, 0.3, name, 14, WHITE, True)
        add_text(slide, 9.98, yy + 0.35, 2.1, 0.28, detail, 11, "C9D8E8")
    add_text(slide, 9.12, 5.38, 3.0, 0.5, "下一步：VLM候选决策（1C）", 13, "8FD9D1", True)
    add_text(slide, 0.8, 6.57, 7.5, 0.38,
             "GOPT/EMS-style几何候选  +  VLM高层决策  +  ReMe动态经验记忆", 14, WHITE, True)
    add_text(slide, 12.0, 7.12, 0.65, 0.23, "1/10", 9, "AFC2D4", True, PP_ALIGN.RIGHT)

    # 2. 技术路线
    slide = base_slide(prs, 2, "课题总体技术路线与当前边界",
                       "前期先验证VLM、规划器和经验记忆；Gazebo与Aubo i5放在最后做闭环验证")
    flow = [
        ("数据集", "真实尺寸/重量\n语义属性/规则", TEAL, PALE_TEAL),
        ("几何规划器", "EMS-style候选\n硬约束过滤", BLUE, PALE_BLUE),
        ("VLM决策", "理解指令\n从Top-K选ID", AMBER, PALE_AMBER),
        ("ReMe记忆", "经验检索\n失败后重写", RED, PALE_RED),
        ("机器人验证", "Gazebo → RGB-D\nMoveIt 2 → Aubo", NAVY_2, "E9EEF3"),
    ]
    for i, args in enumerate(flow):
        x = 0.66 + i * 2.52
        add_flow_box(slide, x, 1.68, 2.02, 1.48, *args)
        if i < len(flow) - 1:
            add_arrow(slide, x + 2.08, 2.22, 0.35, 0.35, MUTED)
    add_rect(slide, 0.7, 3.63, 7.55, 2.62, WHITE, LINE)
    add_text(slide, 1.0, 3.9, 3.0, 0.38, "本次汇报覆盖", 18, NAVY, True)
    add_lines(slide, 1.0, 4.4, 6.75, 1.45, [
        "阶段0：最小完整数据集与接口验收",
        "阶段1A：Repository / Environment / Selector / Runner / Evaluator",
        "阶段1B：在线几何候选、传统基线、批量实验、3D动画",
    ], 14, INK, True, 7)
    add_rect(slide, 8.55, 3.63, 4.08, 2.62, NAVY, NAVY)
    add_text(slide, 8.9, 3.94, 3.3, 0.38, "核心分工", 18, WHITE, True)
    add_lines(slide, 8.92, 4.42, 3.25, 1.5, [
        ("规划器：保证物理可行", "C9D8E8", True),
        ("VLM：处理自然语言业务规则", "8FD9D1", True),
        ("ReMe：沉淀可复用程序性经验", "FFD692", True),
    ], 14, WHITE, False, 8)

    # 3. 数据集
    slide = base_slide(prs, 3, "数据集如何构建：真实订单骨架 + 可控语义层",
                       "MVD V0.1用于接口与功能实验；真实工业结论需要后续扩充")
    add_flow_box(slide, 0.7, 1.5, 3.5, 1.28, "BED-BPP", "箱体真实尺寸、重量、订单组成、到货顺序", BLUE, PALE_BLUE)
    add_flow_box(slide, 4.9, 1.5, 3.5, 1.28, "MixedPalletBoxes思想", "易碎、可堆叠、最大承重等语义属性", TEAL, PALE_TEAL)
    add_arrow(slide, 4.35, 1.96, 0.4, 0.35, MUTED)
    add_arrow(slide, 8.55, 1.96, 0.4, 0.35, MUTED)
    add_flow_box(slide, 9.1, 1.5, 3.5, 1.28, "MVD V0.1", "多层3D放置、业务指令、ID/OOD与视觉输入", AMBER, PALE_AMBER)
    metrics = [
        (str(summary["base_scenario_count"]), "基础Scenario", BLUE),
        (str(summary["visible_scenario_file_count"]), "可见Scenario文件", TEAL),
        (str(summary["decision_case_count"]), "固定Decision Case", AMBER),
        (str(stage0["counts"]["checked_candidates"]), "Top-K候选", RED),
    ]
    for i, (v, l, c) in enumerate(metrics):
        add_metric(slide, 0.7 + i * 3.03, 3.15, 2.7, v, l, c)
    add_rect(slide, 0.7, 4.55, 7.15, 1.72, WHITE, LINE)
    add_text(slide, 0.95, 4.78, 2.6, 0.32, "每个箱体的两类信息", 16, NAVY, True)
    add_lines(slide, 0.95, 5.18, 6.45, 0.85, [
        "几何/物理：length、width、height、mass、pose、support、load",
        "语义/业务：fragile、stackable、weight_class、product_category、自然语言规则",
    ], 13.5, INK, True, 6)
    add_rect(slide, 8.15, 4.55, 4.45, 1.72, PALE_AMBER, AMBER)
    add_text(slide, 8.42, 4.78, 3.8, 0.32, "固定三维码垛空间", 16, AMBER, True)
    add_text(slide, 8.42, 5.18, 3.65, 0.78,
             "1200 × 1000 × 850 mm\n坐标原点：容器底面中心\nz_base：箱体底面相对托盘底面的高度", 13, INK)

    # 4. 目录架构
    slide = base_slide(prs, 4, "数据集目录架构：每个文件夹承担一种实验职责")
    add_rect(slide, 0.68, 1.42, 4.05, 5.42, NAVY, NAVY)
    tree = [
        "semantic_pallet_dataset_v1/",
        "├─ catalog/        物料与几何目录",
        "├─ tasks/          场景与固定决策题",
        "├─ candidates/     候选集合与生成审计",
        "├─ annotations/    规则、Oracle与参考计划",
        "├─ splits/         train / val / ID / OOD",
        "├─ streams/        连续任务与规则变化",
        "├─ views/          VLM请求与渲染图",
        "├─ schemas/        JSON格式约束",
        "├─ statistics/     验收与统计证据",
        "└─ scripts/tests/   构建、检查与测试",
    ]
    add_lines(slide, 0.98, 1.72, 3.45, 4.85, tree, 12.5, WHITE, False, 7, MONO)
    groups = [
        ("任务定义", "tasks + catalog", "回答“要放哪些箱、当前放到哪一步”", BLUE, PALE_BLUE),
        ("动作空间", "candidates", "回答“当前物理可行的放置候选有哪些”", TEAL, PALE_TEAL),
        ("隐藏答案", "annotations", "动作后评价；普通VLM输入不可读取", RED, PALE_RED),
        ("公平评测", "splits + streams", "区分训练、验证、ID/OOD和连续记忆任务", AMBER, PALE_AMBER),
        ("可复现证据", "schemas + statistics", "Schema、校验和、独立几何复核、报告", NAVY_2, "E9EEF3"),
    ]
    for i, (title, tag, detail, color, pale) in enumerate(groups):
        yy = 1.42 + i * 1.08
        add_rect(slide, 5.05, yy, 7.55, 0.88, pale, color)
        add_text(slide, 5.28, yy + 0.16, 1.25, 0.3, title, 14, color, True)
        add_text(slide, 6.58, yy + 0.16, 1.52, 0.3, tag, 11, color, True, font=MONO)
        add_text(slide, 8.12, yy + 0.14, 4.05, 0.4, detail, 12.5, INK)

    # 5. JSON关系
    slide = base_slide(prs, 5, "JSON不是孤立文件：它们组成一条可追溯决策链")
    chain = [
        ("SCN_0001", "完整订单\n容器/物料/到货顺序", BLUE, PALE_BLUE),
        ("DC_00001", "某一步状态\nplaced + available", TEAL, PALE_TEAL),
        ("CSET_00001", "Top-K候选\n位姿/支撑/几何分", AMBER, PALE_AMBER),
        ("Oracle", "隐藏语义评价\n最优候选ID", RED, PALE_RED),
        ("step事件", "请求/响应/前后状态\n指标/耗时/错误", NAVY_2, "E9EEF3"),
    ]
    for i, args in enumerate(chain):
        x = 0.66 + i * 2.52
        add_flow_box(slide, x, 1.42, 2.02, 1.3, *args)
        if i < 4: add_arrow(slide, x + 2.08, 1.9, 0.35, 0.3, MUTED)
    add_rect(slide, 0.7, 3.08, 5.85, 3.27, WHITE, LINE)
    add_text(slide, 0.97, 3.33, 2.6, 0.32, "关键字段怎么读", 17, NAVY, True)
    fields = [
        ("scenario_id", "一条完整码垛任务"),
        ("decision_case_id", "可独立比较方法的一道单步题"),
        ("candidate_id", "规划器生成的合法动作ID"),
        ("pose.x/y/z_base", "箱体底面中心位置；单位mm"),
        ("state_hash", "绑定当前状态，防止过期响应"),
        ("hard_violations", "独立复核发现的违规列表"),
    ]
    for i, (key, meaning) in enumerate(fields):
        yy = 3.82 + i * 0.38
        add_text(slide, 0.98, yy, 1.72, 0.25, key, 10.5, BLUE, True, font=MONO)
        add_text(slide, 2.74, yy, 3.42, 0.25, meaning, 11.3, INK)
    code = '''{
  "request_id": "DC_00001",
  "state_hash": "848195...",
  "instruction_text": "电子产品箱不得...",
  "candidates": [
    {"display_label":"C01",
     "candidate_id":"cand_af7d...",
     "pose":{"x_mm":-85,"y_mm":-300,
             "z_base_mm":0,"yaw_deg":90}}
  ]
}'''
    add_rect(slide, 6.85, 3.08, 5.75, 3.27, NAVY, NAVY)
    add_text(slide, 7.14, 3.32, 5.15, 2.72, code, 11.5, "DDE8F2", False, font=MONO)
    add_text(slide, 7.14, 6.02, 5.05, 0.22, "同一ID贯穿请求、选择、执行、评价和回放。", 10.5, "8FD9D1", True)

    # 6. VLM决策输入
    slide = base_slide(prs, 6, "一次VLM决策到底看到什么？",
                       "公开请求用于选择；Oracle与结构化真值规则只在选择后评价")
    workspace = DATASET / "views/rendered/DC_00001/workspace.png"
    montage = DATASET / "views/rendered/DC_00001/candidates_montage.png"
    add_picture_crop(slide, workspace, 0.68, 1.42, 4.0, 2.65)
    add_picture_crop(slide, montage, 4.9, 1.42, 4.78, 2.65)
    add_rect(slide, 9.92, 1.42, 2.7, 2.65, PALE_TEAL, TEAL)
    add_text(slide, 10.18, 1.68, 2.18, 0.32, "VLM输出", 17, TEAL, True, PP_ALIGN.CENTER)
    add_text(slide, 10.18, 2.18, 2.18, 1.1,
             "只返回\ncandidate_id\n+ confidence / reason", 15, INK, True, PP_ALIGN.CENTER)
    add_text(slide, 10.18, 3.48, 2.18, 0.27, "不能生成新坐标", 11, RED, True, PP_ALIGN.CENTER)
    add_text(slide, 0.76, 4.2, 3.7, 0.25, "① 当前物料 + 已放状态", 12.5, NAVY, True, PP_ALIGN.CENTER)
    add_text(slide, 5.0, 4.2, 4.55, 0.25, "② Top-K候选图，与C01…CK严格对应", 12.5, NAVY, True, PP_ALIGN.CENTER)
    add_rect(slide, 0.68, 4.7, 5.85, 1.52, PALE_BLUE, BLUE)
    add_text(slide, 0.98, 4.93, 2.0, 0.3, "允许看到", 15, BLUE, True)
    add_text(slide, 0.98, 5.34, 5.05, 0.58,
             "自然语言指令、容器、当前箱体属性、已放箱体、候选位姿、支撑/承重与几何分项、决策前图像", 12.3, INK)
    add_rect(slide, 6.78, 4.7, 5.84, 1.52, PALE_RED, RED)
    add_text(slide, 7.08, 4.93, 2.0, 0.3, "禁止看到", 15, RED, True)
    add_text(slide, 7.08, 5.34, 5.05, 0.58,
             "oracle_candidate_id、candidate_evaluations、ground_truth_rule、split、未来到货物料、选择后反馈", 12.3, INK)

    # 7. 1A
    slide = base_slide(prs, 7, "阶段1A：把演示脚本改造成统一、可替换的实验平台")
    modules = [
        ("Repository", "只读加载\nSCN/DC/CSET", BLUE, PALE_BLUE),
        ("Environment", "reset/observe/step\n状态与哈希", TEAL, PALE_TEAL),
        ("Selector", "统一ID响应\n方法可插拔", AMBER, PALE_AMBER),
        ("Evaluator", "几何/语义\n动作后评分", RED, PALE_RED),
        ("Logger", "JSONL/CSV\n重放与复现", NAVY_2, "E9EEF3"),
    ]
    for i, args in enumerate(modules):
        x = 0.66 + i * 2.52
        add_flow_box(slide, x, 1.42, 2.02, 1.22, *args)
        if i < 4: add_arrow(slide, x + 2.08, 1.86, 0.35, 0.3, MUTED)
    add_rect(slide, 0.7, 3.02, 5.85, 3.22, WHITE, LINE)
    add_text(slide, 0.98, 3.28, 2.6, 0.35, "steps.jsonl：一行就是一次决策", 17, NAVY, True)
    add_lines(slide, 0.98, 3.8, 5.15, 1.92, [
        "身份：experiment / method / scenario / step / request_id",
        "输入：before + request + candidate_set + state_hash",
        "动作：selector_response + candidate_id + fallback/error",
        "结果：after + metrics + timing + termination_reason",
        "用途：复查数据泄漏、重算指标、恢复轨迹、比较方法",
    ], 12.5, INK, True, 5)
    add_rect(slide, 6.82, 3.02, 5.8, 3.22, NAVY, NAVY)
    add_text(slide, 7.13, 3.28, 2.5, 0.34, "阶段1A正式产出", 17, WHITE, True)
    add_metric(slide, 7.12, 3.86, 2.3, "120", "固定决策记录", TEAL)
    add_metric(slide, 9.72, 3.86, 2.3, "1", "在线验收episode", AMBER)
    add_text(slide, 7.15, 5.2, 4.9, 0.7,
             "Random / Geometry-Greedy / Handcrafted-Rule / Candidate Oracle\n同一CSET配对比较；Oracle命中100%；两次运行核心哈希一致", 12.3, "DDE8F2")

    # 8. 1B
    slide = base_slide(prs, 8, "阶段1B：EMS-style候选生成、在线规划与Top-K冻结")
    pipeline = ["极值点/支撑层", "0°/90°姿态", "硬约束过滤", "几何评分", "多样性Top-K"]
    colors = [BLUE, TEAL, RED, AMBER, NAVY_2]
    pales = [PALE_BLUE, PALE_TEAL, PALE_RED, PALE_AMBER, "E9EEF3"]
    for i, label in enumerate(pipeline):
        x = 0.66 + i * 2.52
        add_rect(slide, x, 1.38, 2.02, 0.75, pales[i], colors[i])
        add_text(slide, x + 0.08, 1.59, 1.86, 0.28, label, 13.3, colors[i], True, PP_ALIGN.CENTER)
        if i < 4: add_arrow(slide, x + 2.08, 1.58, 0.35, 0.3, MUTED)
    add_rect(slide, 0.7, 2.48, 4.35, 3.78, WHITE, LINE)
    add_text(slide, 0.98, 2.74, 2.7, 0.35, "物理硬约束", 17, NAVY, True)
    add_lines(slide, 0.98, 3.22, 3.72, 1.8, [
        "容器边界与最大高度",
        "三维AABB不重叠",
        "支撑面积比 ≥ 0.70",
        "重心投影在有效支撑区",
        "不可堆叠与传播承重限制",
    ], 12.4, INK, True, 5)
    add_text(slide, 0.98, 5.38, 3.68, 0.55,
             "语义偏好不进入硬过滤，保留给VLM选择。", 12.5, RED, True)
    add_picture_crop(slide, EXP_1B / "figures/topk_sensitivity.png", 5.3, 2.48, 7.3, 3.78)
    add_text(slide, 5.62, 5.91, 6.65, 0.24,
             f"仅在validation上比较K=4/8/16/32 → 冻结K={topk['frozen_top_k']}，后续1C沿用", 11.5, NAVY, True, PP_ALIGN.CENTER)

    # 9. 结果
    slide = base_slide(prs, 9, "阶段1B实验结果：几何方法已能被稳定地区分和复现")
    add_picture_crop(slide, EXP_1B / "figures/geometry_comparison.png", 0.68, 1.36, 7.82, 3.38)
    add_metric(slide, 8.78, 1.42, 1.75, "100/100", "Geometry完成", BLUE)
    add_metric(slide, 10.72, 1.42, 1.75, "57/100", "Lowest完成", TEAL)
    add_metric(slide, 8.78, 2.72, 1.75, "50/100", "Pareto完成", AMBER)
    add_metric(slide, 10.72, 2.72, 1.75, "0", "internal error", RED)
    add_rect(slide, 0.68, 5.02, 12.0, 1.25, NAVY, NAVY)
    checks = [
        ("1440", "固定决策"), ("300", "在线episode"),
        ("5124", "缓存候选复核"), (str(replay["returned_candidates_checked"]), "在线Top-K复核"),
        ("0", "物理违规/泄漏/fallback"),
    ]
    for i, (v, label) in enumerate(checks):
        x = 0.96 + i * 2.34
        add_text(slide, x, 5.23, 1.92, 0.36, v, 20, "8FD9D1", True, PP_ALIGN.CENTER)
        add_text(slide, x, 5.68, 1.92, 0.25, label, 10.4, "DDE8F2", False, PP_ALIGN.CENTER)
    add_text(slide, 0.82, 6.47, 11.7, 0.36,
             "解释边界：Geometry-Greedy的100%完成率来自当前筛选后的最小数据集，不能外推为任意工业订单成功率。", 11.2, RED, True, PP_ALIGN.CENTER)

    # 10. 动画与下一步
    slide = base_slide(prs, 10, "产出动画怎么读，以及它如何服务下一阶段")
    gif = EXP_1B / "episodes/geometry_greedy/SCN_0051/animation.gif"
    # PowerPoint放映模式会播放GIF；编辑视图显示首帧。
    add_picture_crop(slide, gif, 0.68, 1.4, 7.0, 4.36)
    add_rect(slide, 7.95, 1.4, 4.7, 2.33, WHITE, LINE)
    add_text(slide, 8.23, 1.65, 3.9, 0.32, "每一步动画表达什么", 17, NAVY, True)
    add_lines(slide, 8.23, 2.1, 3.95, 1.25, [
        "左：当前取料箱及尺寸/重量/属性",
        "中：动作前托盘状态与选中位姿",
        "右：Top-K候选，C01…CK对应JSON",
        "before → selected → 下一步；末尾为最终布局",
    ], 12, INK, True, 4)
    add_rect(slide, 7.95, 3.98, 4.7, 1.78, PALE_AMBER, AMBER)
    add_text(slide, 8.23, 4.2, 3.9, 0.32, "使用边界", 16, AMBER, True)
    add_text(slide, 8.23, 4.62, 3.95, 0.82,
             "这是确定性3D规划可视化，用来验证决策链和展示布局；它不是Gazebo物理仿真，也不代表机械臂执行成功。", 12.1, INK)
    add_rect(slide, 0.68, 6.0, 12.0, 0.76, NAVY, NAVY)
    add_text(slide, 0.98, 6.18, 1.4, 0.3, "下一阶段 1C", 14, "8FD9D1", True)
    add_text(slide, 2.35, 6.16, 9.9, 0.32,
             "固定Top-K=16 → 接入VLM Selector → 只返回candidate_id → 与Geometry/Rule/Oracle配对比较 → 为ReMe积累决策日志", 13, WHITE, True)
    add_text(slide, 0.83, 6.82, 11.8, 0.22,
             "放映提示：左侧GIF在PowerPoint放映模式下自动播放；编辑模式通常只显示首帧。", 9.5, MUTED, False, PP_ALIGN.CENTER)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
