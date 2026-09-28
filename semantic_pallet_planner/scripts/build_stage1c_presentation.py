#!/usr/bin/env python3
"""生成阶段1C十页中文汇报PPT。"""

from pathlib import Path
import json

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
DATASET = ROOT / "semantic_pallet_dataset_v1"
PLANNER = ROOT / "semantic_pallet_planner"
EXP_1B = PLANNER / "outputs/EXP_1B_20260912_FINAL"
EXP_1C = PLANNER / "outputs/EXP_1C_DEEPSEEK_FLASH_FINAL_V1"
CASE = EXP_1C / "selector_inputs/vlm_visual/DC_00157"
OUTPUT = ROOT / "docs/阶段1C项目进展与实验汇报.pptx"

NAVY = "102A43"
NAVY_2 = "183B5B"
BLUE = "2F6BFF"
TEAL = "16A89A"
AMBER = "F2A93B"
RED = "E35D6A"
INK = "172B4D"
MUTED = "52677D"
LIGHT = "F4F7FB"
LINE = "D8E2EC"
WHITE = "FFFFFF"
PALE_BLUE = "EAF1FF"
PALE_TEAL = "E7F7F4"
PALE_AMBER = "FFF4DE"
PALE_RED = "FDECEE"
FONT = "Noto Sans CJK SC"
MONO = "Noto Sans Mono CJK SC"


def rgb(value):
    return RGBColor.from_string(value)


def rect(slide, x, y, w, h, fill, line=None, rounded=True):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line or fill)
    if rounded:
        try:
            shape.adjustments[0] = 0.08
        except Exception:
            pass
    return shape


def text(slide, x, y, w, h, value, size=16, color=INK, bold=False,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, font=FONT,
         margin=0.04, line_spacing=1.04):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = value
    p.alignment = align
    p.line_spacing = line_spacing
    for run in p.runs:
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
    return box


def lines(slide, x, y, w, h, items, size=13, color=INK, bullet=False,
          gap=4, font=FONT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    for index, item in enumerate(items):
        if isinstance(item, tuple):
            value, item_color, item_bold = item
        else:
            value, item_color, item_bold = item, color, False
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.text = ("• " if bullet else "") + value
        p.space_after = Pt(gap)
        p.line_spacing = 1.08
        for run in p.runs:
            run.font.name = font
            run.font.size = Pt(size)
            run.font.bold = item_bold
            run.font.color.rgb = rgb(item_color)
    return box


def title(slide, page, heading, subtitle=None):
    text(slide, 0.65, 0.34, 0.55, 0.42, f"{page:02d}", 15, BLUE, True,
         valign=MSO_ANCHOR.MIDDLE)
    text(slide, 1.18, 0.25, 11.45, 0.58, heading, 25, NAVY, True,
         valign=MSO_ANCHOR.MIDDLE)
    if subtitle:
        text(slide, 1.2, 0.80, 11.1, 0.30, subtitle, 10.5, MUTED)
    rect(slide, 0.65, 1.10, 12.0, 0.025, BLUE, rounded=False)


def footer(slide, page):
    text(slide, 0.68, 7.17, 8.0, 0.18,
         "Semantic Pallet Planner · 阶段1C · DeepSeek Flash", 8.5, MUTED)
    text(slide, 12.0, 7.12, 0.65, 0.23, f"{page}/10", 9, MUTED, True,
         PP_ALIGN.RIGHT)


def base(prs, page, heading, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(LIGHT)
    title(slide, page, heading, subtitle)
    footer(slide, page)
    return slide


def flow_box(slide, x, y, w, h, heading, detail, color, pale):
    rect(slide, x, y, w, h, pale, color)
    text(slide, x + 0.12, y + 0.13, w - 0.24, 0.30, heading, 14.5, color, True,
         PP_ALIGN.CENTER)
    text(slide, x + 0.10, y + 0.50, w - 0.20, h - 0.57, detail, 10.3, MUTED,
         align=PP_ALIGN.CENTER)


def arrow(slide, x, y, w=0.34, h=0.30, color=MUTED):
    shape = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    shape.line.color.rgb = rgb(color)
    return shape


def metric(slide, x, y, w, value, label, color=BLUE):
    rect(slide, x, y, w, 1.02, WHITE, LINE)
    rect(slide, x, y, 0.07, 1.02, color, rounded=False)
    text(slide, x + 0.19, y + 0.10, w - 0.28, 0.40, value, 23, color, True)
    text(slide, x + 0.19, y + 0.57, w - 0.28, 0.25, label, 10.3, MUTED)


def picture(slide, path, x, y, w, h, line=LINE):
    path = Path(path)
    with Image.open(path) as image:
        iw, ih = image.size
    target = w / h
    source = iw / ih
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    if source > target:
        visible = target / source
        pic.crop_left = pic.crop_right = (1 - visible) / 2
    elif source < target:
        visible = source / target
        pic.crop_top = pic.crop_bottom = (1 - visible) / 2
    border = rect(slide, x, y, w, h, WHITE, line, rounded=False)
    border.fill.background()
    return pic


def picture_contain(slide, path, x, y, w, h, line=LINE):
    """完整显示图片，不裁掉纵向候选网格。"""
    path = Path(path)
    with Image.open(path) as image:
        iw, ih = image.size
    source = iw / ih
    target = w / h
    if source > target:
        draw_w, draw_h = w, w / source
        draw_x, draw_y = x, y + (h - draw_h) / 2
    else:
        draw_w, draw_h = h * source, h
        draw_x, draw_y = x + (w - draw_w) / 2, y
    rect(slide, x, y, w, h, WHITE, line, rounded=False)
    return slide.shapes.add_picture(str(path), Inches(draw_x), Inches(draw_y),
                                    Inches(draw_w), Inches(draw_h))


def code_box(slide, x, y, w, h, heading, content, accent=BLUE, size=10.2):
    rect(slide, x, y, w, h, NAVY, NAVY)
    rect(slide, x, y, w, 0.34, accent, accent, rounded=False)
    text(slide, x + 0.13, y + 0.04, w - 0.26, 0.23, heading, 10.5, WHITE, True,
         font=MONO)
    text(slide, x + 0.16, y + 0.48, w - 0.32, h - 0.60, content, size, "D9E6F2",
         font=MONO, line_spacing=1.0)


def add_table(slide, x, y, widths, row_h, headers, rows, colors=None, font_size=10.5):
    total = sum(widths)
    rect(slide, x, y, total, row_h, NAVY, NAVY, rounded=False)
    cursor = x
    for width, value in zip(widths, headers):
        text(slide, cursor + 0.04, y + 0.06, width - 0.08, row_h - 0.10, value,
             font_size, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        cursor += width
    for r_index, row in enumerate(rows):
        yy = y + row_h * (r_index + 1)
        fill = WHITE if r_index % 2 == 0 else "EEF3F8"
        rect(slide, x, yy, total, row_h, fill, LINE, rounded=False)
        cursor = x
        for c_index, (width, value) in enumerate(zip(widths, row)):
            color = colors[c_index] if colors else INK
            text(slide, cursor + 0.04, yy + 0.04, width - 0.08, row_h - 0.08, str(value),
                 font_size, color, c_index == 0, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
            cursor += width


def build():
    dataset = json.loads((DATASET / "statistics/dataset_summary.json").read_text())
    acceptance = json.loads((EXP_1C / "acceptance_report.json").read_text())
    protocol = json.loads((EXP_1C / "metrics/protocol_summary.json").read_text())
    paired = json.loads((EXP_1C / "metrics/paired_comparisons.json").read_text())
    cost = json.loads((EXP_1C / "metrics/cost_summary.json").read_text())

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1. 封面
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(NAVY)
    rect(slide, 0, 0, 0.16, 7.5, TEAL, rounded=False)
    text(slide, 0.78, 0.72, 7.0, 0.35, "硕士课题阶段汇报 · 2026.09", 13, "8FD9D1", True)
    text(slide, 0.75, 1.25, 8.3, 1.42,
         "阶段1C：VLM语义候选决策\n项目进展与正式实验", 29, WHITE, True,
         valign=MSO_ANCHOR.MIDDLE)
    text(slide, 0.80, 2.90, 7.7, 0.48,
         "数据逻辑 · 几何候选算法 · DeepSeek交互 · 全量实验结果", 16, "C9D8E8")
    rect(slide, 8.78, 0.78, 3.78, 5.62, NAVY_2, "315A78")
    text(slide, 9.10, 1.08, 3.15, 0.40, "当前完成状态", 18, WHITE, True)
    stages = [
        ("阶段0", "数据与接口", TEAL, "PASS"),
        ("阶段1A", "统一实验平台", BLUE, "PASS"),
        ("阶段1B", "几何规划与动画", AMBER, "PASS"),
        ("阶段1C", "VLM固定候选实验", RED, "AUTO PASS"),
    ]
    for i, (name, detail, color, status) in enumerate(stages):
        yy = 1.70 + i * 0.99
        rect(slide, 9.08, yy, 0.66, 0.66, color, color)
        text(slide, 9.08, yy + 0.09, 0.66, 0.34, "✓", 21, WHITE, True, PP_ALIGN.CENTER)
        text(slide, 9.92, yy - 0.01, 1.15, 0.28, name, 13.5, WHITE, True)
        text(slide, 9.92, yy + 0.30, 1.95, 0.25, detail, 10.2, "C9D8E8")
        text(slide, 11.20, yy + 0.04, 1.00, 0.24, status, 9.2, "8FD9D1", True,
             PP_ALIGN.RIGHT)
    text(slide, 9.10, 5.80, 3.05, 0.34, "下一步：阶段1D在线闭环", 12.5, "8FD9D1", True)
    text(slide, 0.80, 6.55, 7.9, 0.38,
         "确定性物理安全边界 + VLM语义选择 + 可审计实验链路", 14, WHITE, True)
    text(slide, 12.0, 7.12, 0.65, 0.23, "1/10", 9, "AFC2D4", True, PP_ALIGN.RIGHT)

    # 2. 进展全景
    slide = base(prs, 2, "项目进展全景：工程链路已跑通，研究结论仍需稳健性验证")
    stage_cards = [
        ("阶段0", "MVD-v0.1\nSchema / 候选 / Oracle", TEAL, PALE_TEAL),
        ("阶段1A", "Repository / Runner\n协议 / 日志 / 回放", BLUE, PALE_BLUE),
        ("阶段1B", "EMS-style候选\nTop-K / 在线episode", AMBER, PALE_AMBER),
        ("阶段1C", "DeepSeek Text/Visual\n360题正式矩阵", RED, PALE_RED),
    ]
    for i, card in enumerate(stage_cards):
        x = 0.72 + i * 3.05
        flow_box(slide, x, 1.48, 2.58, 1.34, *card)
        if i < 3:
            arrow(slide, x + 2.68, 1.98)
    metric(slide, 0.72, 3.22, 2.70, "360", "固定 Decision Case", BLUE)
    metric(slide, 3.63, 3.22, 2.70, "2160", "六方法决策记录", TEAL)
    metric(slide, 6.54, 3.22, 2.70, "720", "真实 VLM 输入", AMBER)
    metric(slide, 9.45, 3.22, 2.70, "PASS", "阶段1C自动验收", RED)
    rect(slide, 0.72, 4.58, 7.35, 1.66, WHITE, LINE)
    text(slide, 0.98, 4.82, 3.2, 0.32, "已完成", 16.5, NAVY, True)
    lines(slide, 0.98, 5.20, 6.72, 0.80, [
        "全量六方法公平比较；候选映射、缓存、回退和成本指标完整",
        "独立物理回放、模型输入泄漏扫描、源码/数据/结果哈希均留档",
    ], 12.5, INK, True, 6)
    rect(slide, 8.35, 4.58, 4.25, 1.66, PALE_AMBER, AMBER)
    text(slide, 8.62, 4.82, 3.4, 0.32, "尚需完成", 16.5, AMBER, True)
    lines(slide, 8.62, 5.20, 3.55, 0.83, [
        "研究者人工签认与位置稳健性审计",
        "重复调用、在线VLM episode、ReMe记忆",
    ], 12.2, INK, True, 6)

    # 3. 数据逻辑
    slide = base(prs, 3, "数据逻辑：真实订单骨架 + 确定性合成语义 + 隐藏评价层")
    flow = [
        ("BED-BPP", "尺寸/重量\n订单/到货顺序", BLUE, PALE_BLUE),
        ("属性构建", "材质/易碎\n可堆叠/承重", TEAL, PALE_TEAL),
        ("Scenario", "完整任务\n自然语言规则", AMBER, PALE_AMBER),
        ("Decision Case", "固定状态\n当前箱体", RED, PALE_RED),
        ("Candidate + Oracle", "公开动作空间\n隐藏事后评价", NAVY_2, "E9EEF3"),
    ]
    for i, item in enumerate(flow):
        x = 0.66 + i * 2.52
        flow_box(slide, x, 1.44, 2.02, 1.30, *item)
        if i < 4:
            arrow(slide, x + 2.08, 1.92)
    counts = [
        (str(dataset["base_scenario_count"]), "基础场景", BLUE),
        (str(dataset["visible_scenario_file_count"]), "可见场景文件", TEAL),
        (str(dataset["decision_case_count"]), "固定决策题", AMBER),
        (str(dataset["item_count"]), "物料实例", RED),
    ]
    for i, item in enumerate(counts):
        metric(slide, 0.72 + i * 3.00, 3.10, 2.65, *item)
    rect(slide, 0.72, 4.48, 6.10, 1.75, WHITE, LINE)
    text(slide, 0.98, 4.72, 3.0, 0.30, "公开给 VLM", 16, NAVY, True)
    lines(slide, 0.98, 5.10, 5.48, 0.93, [
        "自然语言指令、当前状态、物料属性、匿名候选位置与几何特征",
        "候选只包含当前决策可知信息；未来物料不参与固定题选择",
    ], 12.3, INK, True, 6)
    rect(slide, 7.10, 4.48, 5.50, 1.75, PALE_RED, RED)
    text(slide, 7.38, 4.72, 3.0, 0.30, "隐藏到选择之后", 16, RED, True)
    lines(slide, 7.38, 5.10, 4.82, 0.93, [
        "结构化规则类型、Oracle候选ID、每候选语义分和数据划分",
        "注意：语义属性与顶部承重仍是合成值，不能代表工业真值",
    ], 12.3, INK, True, 6)

    # 4. 几何算法与动画
    slide = base(prs, 4, "几何算法：先保证物理可行，再把有限候选交给 VLM",
                 "右侧为阶段1B真实在线episode动画；放映时GIF自动播放")
    gif = EXP_1B / "episodes/geometry_greedy/SCN_0051/animation.gif"
    picture(slide, gif, 0.68, 1.40, 7.15, 4.55)
    text(slide, 0.80, 6.04, 6.85, 0.27,
         "SCN_0051 · Geometry-Greedy · 逐箱生成候选、选择、提交并更新状态", 10.5, MUTED)
    steps = [
        ("1", "极值点枚举", "托盘边缘、已放箱边缘、箱顶支撑层"),
        ("2", "姿态展开", "水平 yaw = 0° / 90°"),
        ("3", "硬约束过滤", "边界、碰撞、70%支撑、重心、承重"),
        ("4", "几何评分", "紧凑度35% + 高度25% + 稳定25% + 平衡15%"),
        ("5", "多样性Top-K", "Pareto前沿 + 空间分层，冻结 K=16"),
    ]
    for i, (num, heading, detail) in enumerate(steps):
        yy = 1.42 + i * 0.95
        rect(slide, 8.18, yy, 0.56, 0.56, BLUE, BLUE)
        text(slide, 8.18, yy + 0.07, 0.56, 0.30, num, 15, WHITE, True, PP_ALIGN.CENTER)
        text(slide, 8.92, yy - 0.01, 3.30, 0.28, heading, 13.5, NAVY, True)
        text(slide, 8.92, yy + 0.32, 3.48, 0.38, detail, 10.6, MUTED)
    rect(slide, 8.18, 6.18, 4.22, 0.50, PALE_TEAL, TEAL)
    text(slide, 8.34, 6.29, 3.88, 0.22, "提交前由独立检查器再次复核候选", 11.2, TEAL, True,
         PP_ALIGN.CENTER)

    # 5. 阶段1C核心链路
    slide = base(prs, 5, "阶段1C核心逻辑：模型只在冻结合法候选中做语义选择")
    flow = [
        ("公开请求", "状态 + 指令\n合法Candidate Set", BLUE, PALE_BLUE),
        ("匿名映射", "内部ID删除\n重排为C01…C16", TEAL, PALE_TEAL),
        ("Text / Visual", "结构化JSON\n视觉额外两张图", AMBER, PALE_AMBER),
        ("严格解析", "Schema + request_id\nstate_hash + 合法编号", RED, PALE_RED),
        ("事后评价", "恢复内部ID\n物理复核 + Oracle", NAVY_2, "E9EEF3"),
    ]
    for i, item in enumerate(flow):
        x = 0.66 + i * 2.52
        flow_box(slide, x, 1.48, 2.02, 1.34, *item)
        if i < 4:
            arrow(slide, x + 2.08, 1.98)
    rect(slide, 0.72, 3.22, 7.02, 2.90, WHITE, LINE)
    text(slide, 0.98, 3.48, 3.2, 0.32, "安全与公平机制", 17, NAVY, True)
    mechanisms = [
        "六种方法面对同一冻结 Candidate Set，选择不会改变下一道固定题",
        "模型不生成坐标，无法绕过物理合法候选空间",
        "候选显示顺序按题目和种子确定性随机化，映射单独留审计文件",
        "最多一次格式修复；仍失败则按冻结策略回退 Geometry-Greedy",
        "选择完成后才读取隐藏规则和 Oracle，计算 utility / regret / 违规",
    ]
    lines(slide, 0.98, 3.94, 6.40, 1.85, mechanisms, 12.0, INK, True, 6)
    rect(slide, 8.02, 3.22, 4.58, 2.90, NAVY, NAVY)
    text(slide, 8.34, 3.50, 3.88, 0.32, "响应必须满足", 17, WHITE, True)
    lines(slide, 8.34, 3.98, 3.82, 1.62, [
        "schema_version 正确",
        "request_id 与当前题一致",
        "state_hash 与当前状态一致",
        "display_candidate_id 属于 C01…CK",
        "禁止坐标字段和额外顶层字段",
    ], 12.2, "D9E6F2", True, 7)

    # 6. 真实视觉输入
    slide = base(prs, 6, "模型实际看到什么：结构化状态 + 工作区图 + 候选拼图",
                 "示例来自正式实验 EXP_1C_DEEPSEEK_FLASH_FINAL_V1 / DC_00157")
    picture(slide, CASE / "workspace.png", 0.68, 1.38, 7.06, 3.22)
    picture_contain(slide, CASE / "candidates_montage.png", 8.02, 1.38, 4.60, 5.02)
    text(slide, 0.78, 4.72, 6.82, 0.25, "工作区图：取料区 + 当前托盘状态", 10.5, MUTED,
         align=PP_ALIGN.CENTER)
    text(slide, 8.12, 6.47, 4.40, 0.25, "候选拼图：完整 C01～C16", 10.5, MUTED,
         align=PP_ALIGN.CENTER)
    rect(slide, 0.68, 5.18, 7.06, 1.22, PALE_AMBER, AMBER)
    text(slide, 0.92, 5.40, 6.58, 0.72,
         "Visual 模式并非让模型自由定位：图片用于辅助理解空间关系，最终仍只能返回一个 C 编号。\n"
         "当前图片是确定性仿真渲染，不是真实相机 RGB-D；因此视觉收益必须单独验证。",
         11.5, INK, True, PP_ALIGN.CENTER)

    # 7. JSON证据
    slide = base(prs, 7, "关键 JSON 证据：请求、响应、内部映射与验收均可追溯",
                 "以下内容为真实产物的精简摘录，文本框可在PowerPoint中直接编辑")
    request_json = '''{
  "request_id": "DC_00157",
  "instruction_text":
    "易碎箱应尽量不承载其他箱体，并优先放在较高层。",
  "available_items": [{
    "item_id": "item_02", "fragile": true,
    "weight_class": "heavy", "stackable": true
  }],
  "candidates": [{
    "display_label": "C06",
    "pose": {"x_mm": 400, "y_mm": -350,
             "z_base_mm": 0, "yaw_deg": 0}
  }]
}'''
    response_json = '''{
  "schema_version": "pallet_vlm_choice_v1",
  "request_id": "DC_00157",
  "state_hash": "0fc4aa2d...1423a",
  "display_candidate_id": "C06",
  "confidence": 0.72,
  "applied_rule": "易碎箱优先高层且不承载"
}

本地映射：
C06 → cand_13b31ca377ae'''
    acceptance_json = '''{
  "stage": "1c",
  "status": "PASS",
  "fixed_rows": 2160,
  "decision_cases": 360,
  "checks": [
    "same_frozen_candidate_set",
    "zero_model_input_leakage",
    "zero_physical_hard_violations",
    "vlm_valid_rate >= 98%",
    "replay_pass"
  ]
}'''
    code_box(slide, 0.68, 1.38, 4.02, 4.96, "public_request.json", request_json, BLUE, 9.3)
    code_box(slide, 4.86, 1.38, 3.66, 4.96, "response.json + mapping.audit.json",
             response_json, TEAL, 9.1)
    code_box(slide, 8.68, 1.38, 3.96, 4.96, "acceptance_report.json",
             acceptance_json, RED, 9.3)

    # 8. 正式结果
    slide = base(prs, 8, "正式实验结果：VLM优于纯几何，Visual相对Text增益很小")
    headers = ["方法", "平均utility", "Oracle命中", "业务硬违规"]
    rows = [
        ["Random Valid", "0.670115", "32 / 360", "50"],
        ["Geometry Greedy", "0.637921", "108 / 360", "60"],
        ["Handcrafted Rule*", "0.765367", "360 / 360", "39"],
        ["DeepSeek Text", "0.695990", "58 / 360", "45"],
        ["DeepSeek Visual", "0.698657", "58 / 360", "45"],
        ["Candidate Oracle*", "0.765367", "360 / 360", "39"],
    ]
    add_table(slide, 0.68, 1.42, [2.60, 1.52, 1.58, 1.50], 0.55,
              headers, rows, [NAVY, INK, INK, INK], 10.4)
    text(slide, 0.72, 5.47, 6.98, 0.35,
         "* Handcrafted Rule 与 Candidate Oracle 使用隐藏信息，是特权基线/上界。", 9.5, MUTED)
    rect(slide, 8.14, 1.42, 4.48, 2.10, PALE_BLUE, BLUE)
    text(slide, 8.42, 1.68, 3.88, 0.32, "同题配对比较", 17, BLUE, True)
    p1 = paired["vlm_text_minus_geometry_greedy"]
    p2 = paired["vlm_visual_minus_geometry_greedy"]
    p3 = paired["vlm_visual_minus_vlm_text"]
    lines(slide, 8.42, 2.12, 3.74, 1.10, [
        f"Text − Geometry：+{p1['mean_soft_utility_delta']:.6f}  ·  {p1['wins']}/{p1['ties']}/{p1['losses']}",
        f"Visual − Geometry：+{p2['mean_soft_utility_delta']:.6f}  ·  {p2['wins']}/{p2['ties']}/{p2['losses']}",
        f"Visual − Text：+{p3['mean_soft_utility_delta']:.6f}  ·  {p3['wins']}/{p3['ties']}/{p3['losses']}",
    ], 11.3, INK, False, 7)
    rect(slide, 8.14, 3.78, 4.48, 2.48, WHITE, LINE)
    text(slide, 8.42, 4.04, 3.88, 0.32, "协议与成本", 17, NAVY, True)
    lines(slide, 8.42, 4.48, 3.74, 1.42, [
        f"Text：{protocol['vlm_text']['valid_count']}/360 有效，1次修复，0回退",
        f"Visual：{protocol['vlm_visual']['valid_count']}/360 有效，3次回退",
        f"输入token：Text {cost['vlm_text']['input_tokens_total']:,}",
        f"输入token：Visual {cost['vlm_visual']['input_tokens_total']:,}",
        "费用字段为0仅因YAML未配置单价",
    ], 11.0, INK, True, 6)

    # 9. 错误案例和边界
    slide = base(prs, 9, "案例审计：工程PASS不等于每次语义判断都正确")
    cases = [
        ("DC_00153", "支撑关系反向理解", "VLM 0.267306 / Oracle 0.433653", RED, PALE_RED),
        ("DC_00177", "电子产品与重物接触", "VLM业务硬违规=1；Oracle utility=1.0", AMBER, PALE_AMBER),
        ("DC_00178", "整体质心 ≠ 当前重箱位置", "VLM 0.301935 / Oracle 0.795742", BLUE, PALE_BLUE),
    ]
    for i, (case_id, heading, detail, color, pale) in enumerate(cases):
        yy = 1.42 + i * 1.22
        rect(slide, 0.72, yy, 6.15, 0.96, pale, color)
        text(slide, 0.96, yy + 0.16, 1.20, 0.27, case_id, 13.5, color, True, font=MONO)
        text(slide, 2.18, yy + 0.14, 2.78, 0.30, heading, 13.2, NAVY, True)
        text(slide, 2.18, yy + 0.51, 4.20, 0.25, detail, 10.8, MUTED)
    rect(slide, 0.72, 5.20, 6.15, 1.08, NAVY, NAVY)
    text(slide, 0.96, 5.43, 5.65, 0.56,
         "自动验收关注矩阵完整、协议有效、物理安全、泄漏与回放。\n"
         "业务语义表现属于研究结果，必须通过案例审计和统计单独讨论。",
         12.1, WHITE, True, PP_ALIGN.CENTER)
    picture_contain(slide, CASE / "candidates_montage.png", 7.16, 1.42, 5.45, 3.55)
    rect(slide, 7.16, 5.20, 5.45, 1.08, PALE_AMBER, AMBER)
    lines(slide, 7.42, 5.40, 4.92, 0.64, [
        "当前边界：合成语义、固定单步题、渲染图而非RGB-D",
        "泄漏扫描尚无图片像素检查；future_item_hits仍为占位",
    ], 10.9, INK, True, 5)

    # 10. 结论与下一步
    slide = base(prs, 10, "阶段结论与下一步：从“链路跑通”转向“语义稳健与在线闭环”")
    rect(slide, 0.72, 1.42, 5.72, 4.92, NAVY, NAVY)
    text(slide, 1.02, 1.74, 4.98, 0.36, "当前可以确认", 19, WHITE, True)
    lines(slide, 1.02, 2.26, 4.96, 3.34, [
        "阶段0、1A、1B以及阶段1C自动验收均已通过",
        "几何模块形成确定性物理安全边界，VLM只做离散语义选择",
        "DeepSeek Text/Visual在平均语义效用上均高于Geometry Greedy",
        "Visual相对Text仅+0.002667，当前尚无稳定视觉收益证据",
        "正式实验、请求响应、图片、映射、缓存、回放和哈希均可审计",
    ], 13.0, "D9E6F2", True, 10)
    text(slide, 1.02, 5.80, 4.96, 0.28,
         "结论边界：工程通过 ≠ 工业效果成立", 12.2, "8FD9D1", True,
         PP_ALIGN.CENTER)
    next_steps = [
        ("01", "研究者签认 + 位置偏差审计", BLUE, PALE_BLUE),
        ("02", "在train/validation升级关系特征与Prompt", TEAL, PALE_TEAL),
        ("03", "重复调用，给出稳定性和视觉增益置信区间", AMBER, PALE_AMBER),
        ("04", "冻结选择器，进入阶段1D在线完整episode", RED, PALE_RED),
        ("05", "阶段1E接入ReMe；最后进入RGB-D/Gazebo/机器人", NAVY_2, "E9EEF3"),
    ]
    for i, (num, detail, color, pale) in enumerate(next_steps):
        yy = 1.42 + i * 0.97
        rect(slide, 6.76, yy, 5.84, 0.74, pale, color)
        text(slide, 6.95, yy + 0.15, 0.58, 0.28, num, 13, color, True,
             PP_ALIGN.CENTER)
        text(slide, 7.62, yy + 0.14, 4.62, 0.31, detail, 12.1, INK, True)
    rect(slide, 6.76, 6.38, 5.84, 0.38, PALE_TEAL, TEAL)
    text(slide, 6.92, 6.46, 5.52, 0.19,
         "建议：Text作为下一阶段主配置，Visual保留为消融对照", 10.7, TEAL, True,
         PP_ALIGN.CENTER)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
